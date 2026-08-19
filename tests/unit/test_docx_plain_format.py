"""Round-trip tests for the DOCX plain extractor (Plain Text Mode).

These cover the two historical failure modes:
- tables silently dropped (the old extractor iterated doc.paragraphs only)
- inline formatting flattened (bold/italic/hyperlinks destroyed)
"""
import pytest

docx = pytest.importorskip("docx")

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.core.docx.plain_extractor import (
    build_minimal_docx,
    extract_plain_paragraphs,
)


def _add_hyperlink(paragraph, doc, text, url):
    part = doc.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


@pytest.fixture
def rich_docx(tmp_path):
    doc = Document()
    doc.add_heading("Main title", level=1)

    p = doc.add_paragraph()
    p.add_run("Normal with ")
    r = p.add_run("bold")
    r.bold = True
    p.add_run(" and ")
    r = p.add_run("italic")
    r.italic = True
    p.add_run(" and ")
    r = p.add_run("both")
    r.bold = True
    r.italic = True
    p.add_run(".")

    p2 = doc.add_paragraph()
    p2.add_run("Visit ")
    _add_hyperlink(p2, doc, "the guide", "https://example.com/guide")
    p2.add_run(" today.")

    doc.add_paragraph("Item one", style="List Bullet")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alpha"
    run = table.cell(1, 1).paragraphs[0].add_run("Beta strong")
    run.bold = True

    doc.add_paragraph("Closing paragraph after the table.")

    path = tmp_path / "rich.docx"
    doc.save(str(path))
    return str(path)


def _body_kinds(doc):
    kinds = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            par = Paragraph(child, doc)
            if par.text.strip():
                kinds.append(("p", par))
        elif child.tag == qn("w:tbl"):
            kinds.append(("table", Table(child, doc)))
    return kinds


class TestExtraction:
    def test_blocks_in_document_order_with_styles(self, rich_docx):
        content = extract_plain_paragraphs(rich_docx)
        styles = content.paragraphs_style
        assert styles[0] == "heading1"
        assert "list" in styles
        assert styles.count("table_cell") == 4
        # Closing paragraph must come AFTER the table cells (document order).
        assert styles[-1] == "normal"
        assert content.paragraphs_text[-1].startswith("Closing")

    def test_inline_formatting_encoded_as_markdown(self, rich_docx):
        content = extract_plain_paragraphs(rich_docx)
        body = content.paragraphs_text[1]
        assert "**bold**" in body
        assert "*italic*" in body
        assert "***both***" in body

    def test_hyperlink_encoded_with_url(self, rich_docx):
        content = extract_plain_paragraphs(rich_docx)
        link_par = content.paragraphs_text[2]
        assert "[the guide](https://example.com/guide)" in link_par

    def test_table_geometry_recorded(self, rich_docx):
        content = extract_plain_paragraphs(rich_docx)
        assert len(content.tables) == 1
        spec = next(iter(content.tables.values()))
        assert spec.rows == 2 and spec.cols == 2
        assert len(spec.cell_blocks) == 4
        # Bold cell content carries its markers.
        bold_cell = content.paragraphs_text[spec.grid[1][1]]
        assert bold_cell == "**Beta strong**"


class TestRebuild:
    def test_round_trip_preserves_structure_and_formatting(self, rich_docx, tmp_path):
        content = extract_plain_paragraphs(rich_docx)
        translated = [t.replace("Normal", "Plain") for t in content.paragraphs_text]
        out_path = tmp_path / "out.docx"
        build_minimal_docx(translated, content, str(out_path))

        out = Document(str(out_path))
        kinds = _body_kinds(out)

        tables = [obj for kind, obj in kinds if kind == "table"]
        assert len(tables) == 1
        table = tables[0]
        assert len(table.rows) == 2 and len(table.columns) == 2
        bold_runs = [
            run
            for p in table.cell(1, 1).paragraphs
            for run in p.runs
            if run.text
        ]
        assert bold_runs and bold_runs[0].bold

        # Paragraph after the table survives in position.
        assert kinds[-1][0] == "p"
        assert "Closing" in kinds[-1][1].text

        # Inline formatting decoded back into runs.
        body_par = kinds[1][1]
        flags = {(r.text, bool(r.bold), bool(r.italic)) for r in body_par.runs}
        assert ("bold", True, False) in flags
        assert ("italic", False, True) in flags
        assert ("both", True, True) in flags

    def test_empty_translation_cells_do_not_crash(self, rich_docx, tmp_path):
        content = extract_plain_paragraphs(rich_docx)
        translated = ["" for _ in content.paragraphs_text]
        out_path = tmp_path / "empty.docx"
        build_minimal_docx(translated, content, str(out_path))
        out = Document(str(out_path))
        assert any(
            child.tag == qn("w:tbl") for child in out.element.body.iterchildren()
        )
