from docx import Document
from lxml import etree

from src.core.docx.plain_extractor import extract_plain_paragraphs as extract_docx_plain
from src.core.epub.plain_extractor import extract_plain_paragraphs as extract_epub_plain


def test_epub_plain_extractor_keeps_table_and_figure_text():
    body = etree.fromstring(
        """
        <body xmlns="http://www.w3.org/1999/xhtml">
          <p>Intro text.</p>
          <table><tr><td>Cell A</td><td>Cell B</td></tr></table>
          <figure><figcaption>Figure caption</figcaption></figure>
        </body>
        """
    )

    paragraphs, tags, _images, tables = extract_epub_plain(body)
    joined = "\n".join(paragraphs)

    assert "Intro text." in joined
    assert "Cell A" in joined
    assert "Cell B" in joined
    assert "Figure caption" in joined
    assert "table_cell" in tags
    assert len(tables) == 1


def test_docx_plain_extractor_keeps_table_rows(tmp_path):
    path = tmp_path / "table.docx"
    doc = Document()
    doc.add_paragraph("Intro text.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    doc.save(path)

    content = extract_docx_plain(str(path))

    assert "Intro text." in content.paragraphs_text
    assert "Cell A" in content.paragraphs_text
    assert "Cell B" in content.paragraphs_text
    assert "table_cell" in content.paragraphs_style
    assert len(content.tables) == 1
