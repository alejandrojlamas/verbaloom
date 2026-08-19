from docx import Document
import zipfile

from src.core.layout_sanitizer import (
    sanitize_extracted_pages,
    sanitize_extracted_text,
    should_use_sanitized_text_pipeline,
    strip_inline_style_markers,
)
from src.core.output_formats import extract_readable_text


def test_pdf_layout_sanitizer_removes_running_margins_and_page_numbers():
    pages = [
        "DON QUIJOTE\n1\nTexto real de la primera pagina con-\ntinua aqui.\nEditorial 2026",
        "DON QUIJOTE\n2\nTexto real de la segunda pagina.\nEditorial 2026",
    ]

    cleaned, report = sanitize_extracted_pages(pages, source_type="pdf")

    assert "DON QUIJOTE" not in cleaned
    assert "Editorial 2026" not in cleaned
    assert "\n1\n" not in f"\n{cleaned}\n"
    assert "\n2\n" not in f"\n{cleaned}\n"
    assert "continua aqui" in cleaned
    assert report.removed_repeated_margin_lines >= 2
    assert report.removed_page_artifacts >= 2


def test_text_sanitizer_keeps_content_chapter_numbers_without_page_context():
    cleaned, _report = sanitize_extracted_text(
        "I\n\nEn un lugar de la Mancha.\n\nII\n\nOtra seccion.",
        source_type="pdf",
    )

    assert cleaned.startswith("I")
    assert "\nII\n" in f"\n{cleaned}\n"


def test_epub_sanitizer_removes_standalone_source_links_and_page_numbers():
    cleaned, report = sanitize_extracted_text(
        (
            "Texto real antes del artefacto.\n\n"
            "[OceanofPDF.com]\n"
            "(https://oceanofpdf.com)\n\n"
            "12.\n\n"
            "Texto real despues."
        ),
        source_type="epub",
    )

    assert "Texto real antes" in cleaned
    assert "Texto real despues" in cleaned
    assert "OceanofPDF" not in cleaned
    assert "https://oceanofpdf.com" not in cleaned
    assert "\n12.\n" not in f"\n{cleaned}\n"
    assert report.removed_page_artifacts >= 2


def test_sanitizer_keeps_inline_link_labels_without_urls():
    cleaned, _report = sanitize_extracted_text(
        "Consulta [el apendice](https://example.com/apendice) para contexto narrativo.",
        source_type="epub",
        strip_inline_style=True,
    )

    assert "Consulta el apendice para contexto narrativo." in cleaned
    assert "https://example.com/apendice" not in cleaned


def test_sanitizer_strips_relative_epub_note_link_targets():
    cleaned, _report = sanitize_extracted_text(
        "ver [[23]](../Text/notas.xhtml#nt23) ahora",
        source_type="epub",
        strip_inline_style=True,
    )

    assert cleaned == "ver [23] ahora"
    assert "notas.xhtml" not in cleaned


def test_sanitizer_removes_standalone_raw_url_lines():
    cleaned, _report = sanitize_extracted_text(
        "Texto real.\n\nhttps://oceanofpdf.com\n\n13.\n\nSigue.",
        source_type="epub",
    )

    assert cleaned == "Texto real.\n\nSigue."
    assert "https://oceanofpdf.com" not in cleaned
    assert "\n13.\n" not in f"\n{cleaned}\n"


def test_sanitizer_preserves_legitimate_standalone_publisher_url():
    cleaned, _report = sanitize_extracted_text(
        "Consulte el material complementario.\n\n"
        "https://www.wiley.com/college/block\n\n"
        "Sigue el texto.",
        source_type="epub",
    )

    assert "https://www.wiley.com/college/block" in cleaned


def test_epub_readable_text_applies_source_link_sanitizer(tmp_path):
    epub_path = tmp_path / "source.epub"
    with zipfile.ZipFile(epub_path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr(
            "META-INF/container.xml",
            """<?xml version="1.0"?>
            <container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
              <rootfiles>
                <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
              </rootfiles>
            </container>""",
        )
        zf.writestr(
            "OEBPS/content.opf",
            """<?xml version="1.0"?>
            <package xmlns="http://www.idpf.org/2007/opf" version="3.0">
              <manifest>
                <item id="chap" href="chapter.xhtml" media-type="application/xhtml+xml"/>
              </manifest>
              <spine><itemref idref="chap"/></spine>
            </package>""",
        )
        zf.writestr(
            "OEBPS/chapter.xhtml",
            """<html xmlns="http://www.w3.org/1999/xhtml"><body>
              <p>Texto real antes.</p>
              <p><a href="https://oceanofpdf.com">OceanofPDF.com</a></p>
              <p>12.</p>
              <p>Texto real despues.</p>
            </body></html>""",
        )

    text = extract_readable_text(epub_path)

    assert "Texto real antes." in text
    assert "Texto real despues." in text
    assert "OceanofPDF" not in text
    assert "https://oceanofpdf.com" not in text
    assert "\n12.\n" not in f"\n{text}\n"


def test_strip_inline_style_markers_preserves_literary_asterism():
    cleaned, count = strip_inline_style_markers(
        "**Bold** and *italic* and [guide](https://example.com).\n\n* * *"
    )

    assert cleaned == "Bold and italic and guide.\n\n* * *"
    assert count == 3


def test_docx_readable_text_removes_style_and_keeps_table_content(tmp_path):
    path = tmp_path / "styled.docx"
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "RUNNING HEADER"
    section.footer.paragraphs[0].text = "Page 1"
    doc.add_heading("Main title", level=1)
    p = doc.add_paragraph("A paragraph with ")
    run = p.add_run("bold")
    run.bold = True
    p.add_run(" text.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    doc.save(path)

    text = extract_readable_text(path)

    assert "RUNNING HEADER" not in text
    assert "Page 1" not in text
    assert "**bold**" not in text
    assert "A paragraph with bold text." in text
    assert "Cell A" in text
    assert "Cell B" in text


def test_docx_uses_sanitized_text_pipeline_by_default():
    assert should_use_sanitized_text_pipeline("docx", {}) is True
    assert should_use_sanitized_text_pipeline("docx", {"preserve_source_formatting": True}) is False
    assert should_use_sanitized_text_pipeline("epub", {}) is False
    assert should_use_sanitized_text_pipeline("epub", {"plain_text_mode": True}) is True
