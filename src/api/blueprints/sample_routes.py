"""
Sample & Compare routes.

Runs an arbitrary number of LLM configurations in parallel on N short extracts
of an uploaded book, streaming each cell back to the client over WebSocket as it
completes.
No persistence: state lives in `SampleStateManager` and is dropped after 1
hour or on server restart.
"""
import asyncio
import logging
import os
import random
import threading
import time
import uuid
from pathlib import Path
from typing import TYPE_CHECKING, Any, Dict, List, Optional, Set, Tuple

from flask import Blueprint, jsonify, request

from src.config import (
    MAX_TOKENS_PER_CHUNK, OLLAMA_NUM_CTX,
    NIM_API_ENDPOINT, OPENAI_API_ENDPOINT, OUTPUT_DIR, REQUEST_TIMEOUT,
    SRT_LINES_PER_BLOCK,
)
from src.api.services.path_validator import PathValidator
from src.utils.provider_security import resolve_api_key_for_endpoint
from src.core.book_profiles import build_profile_glossary_block
from src.core.glossary import build_glossary_block, filter_glossary_for_purpose
from src.core.glossary.models import GlossaryConfig
from src.core.llm.factory import create_llm_provider
from src.core.llm.request_deadline import await_llm_call
from src.core.llm_output_guard import guard_llm_output
from src.core.deepseek_pricing import get_deepseek_pricing_status
from src.core.pricing import calculate_usage_cost, get_default_pricing
from src.core.sampling import cap_chunk_text, select_sample_indices
from src.core.text_processor import split_text_into_chunks
from src.prompts.prompts import (
    generate_refinement_prompt, generate_translation_prompt,
)
from src.utils.custom_instructions import is_safe_filename, load_custom_instructions
from src.utils.file_detector import detect_file_type
from src.utils.language_detector import LanguageDetector
from src.utils.text_encoding import clean_text_artifacts

if TYPE_CHECKING:
    from src.api.sample_state import SampleStateManager


# Per-run concurrency cap. The product spec asks for `min(K * N, 8)` to avoid
# hammering providers; this is enforced per sample run via an asyncio.Semaphore.
SAMPLE_CONCURRENCY_CAP = 8

# Sampling defaults used when a request omits the fields (safety net — both
# front-ends always send explicit values). Must mirror the frontend single
# source in src/web/static/js/sample/sample-defaults.js.
DEFAULT_N_SAMPLES = 5
DEFAULT_MAX_CHARS = 400
SAMPLE_REQUEST_TIMEOUT = max(
    30,
    min(REQUEST_TIMEOUT, int(os.getenv("SAMPLE_REQUEST_TIMEOUT", "180"))),
)

logger = logging.getLogger(__name__)


def _clamp_int(value: Any, default: int, lo: int, hi: int) -> int:
    """Parse `value` as int (falling back to `default` when None) and clamp to [lo, hi].

    Raises ValueError/TypeError for non-numeric input so callers can return 400.
    """
    return max(lo, min(hi, int(default if value is None else value)))


def _small_document_warning(total: int, count: int, requested: int) -> Dict[str, Any]:
    """Structured warning for "doc has fewer interior units than requested".

    Returned as {code, params} (not a pre-formatted English string) so the
    client can translate it reactively via the `sample:warning_small_document`
    i18n key, whose {{total}}/{{count}}/{{requested}} placeholders match params.
    """
    return {
        "code": "warning_small_document",
        "params": {"total": total, "count": count, "requested": requested},
    }


def _resolve_api_key(value: Any, env_var_name: str) -> str:
    """Resolve `__USE_ENV__` placeholder to the actual env var value.

    Mirrors `_resolve_api_key` in translation_routes.py — kept inline to avoid
    cross-blueprint imports.
    """
    if value == "__USE_ENV__" or not value:
        return os.getenv(env_var_name, "")
    return value


def _provider_env_var(provider: str) -> str:
    """Return the env var name conventionally used for a provider's API key."""
    return {
        "gemini": "GEMINI_API_KEY",
        "openai": "OPENAI_API_KEY",
        "openrouter": "OPENROUTER_API_KEY",
        "mistral": "MISTRAL_API_KEY",
        "deepseek": "DEEPSEEK_API_KEY",
        "poe": "POE_API_KEY",
        "nim": "NIM_API_KEY",
    }.get(provider.lower(), "")


def _extract_plain_text(file_path: str, file_type: str) -> str:
    """
    Extract the textual content of a file for chunking + sampling.

    For TXT we read directly. For EPUB/DOCX/PDF we reuse plain extractors used
    by Plain Text Mode in the main translate flow. SRT is handled by the
    caller (sampled at the cue-group level, not via this helper).
    """
    ft = file_type.lower()
    if ft == "txt":
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    if ft == "epub":
        return _extract_epub_text(file_path)
    if ft == "docx":
        return _extract_docx_text(file_path)
    if ft == "pdf":
        return _extract_pdf_text(file_path)
    raise ValueError(f"Unsupported file type for sampling: {file_type}")


def _extract_epub_text(file_path: str) -> str:
    """Extract EPUB text in the same spine order used for final TXT export."""
    from src.core.output_formats import extract_readable_text

    return extract_readable_text(file_path)


def _extract_docx_text(file_path: str) -> str:
    """Concatenate paragraph text from a DOCX using the plain extractor."""
    from docx import Document

    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n\n".join(parts)


def _extract_pdf_text(file_path: str) -> str:
    """Extract readable text from a PDF for sampling."""
    from src.core.pdf import extract_pdf_text

    return extract_pdf_text(file_path)


def _load_source_units(file_path: str, file_type: str) -> List[Dict[str, str]]:
    """
    Return a normalized list of "source units" for sampling.

    Every unit is a dict with `main_content` / `context_before` / `context_after`
    so the same item-construction code can serve TXT/EPUB/DOCX/SRT files.
    For TXT/EPUB/DOCX this is just `split_text_into_chunks`; for SRT we group
    cues into blocks of SRT_LINES_PER_BLOCK and synthesize the contexts from
    adjacent blocks.

    Deterministic: identical (file_path, file_type) always returns identical
    units, so an index produced by /initialize remains valid for /extract and
    /run later.
    """
    ft = file_type.lower()
    if ft == "srt":
        from src.core.srt_processor import SRTProcessor

        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

        proc = SRTProcessor()
        subtitles = proc.parse_srt(content)
        if not subtitles:
            raise ValueError("No subtitles found in SRT file")

        block_size = max(1, SRT_LINES_PER_BLOCK)
        blocks: List[str] = []
        for i in range(0, len(subtitles), block_size):
            block_text = "\n".join(
                s["text"] for s in subtitles[i:i + block_size] if s.get("text")
            )
            if block_text.strip():
                blocks.append(block_text)

        total = len(blocks)
        return [
            {
                "main_content": blocks[i],
                "context_before": blocks[i - 1] if i > 0 else "",
                "context_after": blocks[i + 1] if i + 1 < total else "",
            }
            for i in range(total)
        ]

    text = _extract_plain_text(file_path, file_type)
    if not text or not text.strip():
        raise ValueError("File is empty or unreadable")
    return split_text_into_chunks(text, max_tokens_per_chunk=MAX_TOKENS_PER_CHUNK)


def _items_for_indices(
    units: List[Dict[str, str]],
    indices: List[int],
    max_chars: int,
) -> List[Dict[str, Any]]:
    """Build sample items for the given indices, capping each main_content."""
    items: List[Dict[str, Any]] = []
    for idx in indices:
        if idx < 0 or idx >= len(units):
            continue
        unit = units[idx]
        capped, truncated = cap_chunk_text(unit.get("main_content", ""), max_chars)
        items.append({
            "index": idx,
            "source_text": capped,
            "truncated": truncated,
            "context_before": unit.get("context_before", ""),
            "context_after": unit.get("context_after", ""),
        })
    return items


def _build_srt_sample_blocks(file_path: str, n_samples: int, max_chars: int) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """For SRT files, sample N blocks. Returns (items, warnings)."""
    units = _load_source_units(file_path, "srt")
    total = len(units)
    if total < 1:
        raise ValueError("document too small for sampling")

    warnings: List[Dict[str, Any]] = []
    indices = select_sample_indices(total, n_samples)
    if len(indices) < n_samples:
        warnings.append(_small_document_warning(total, len(indices), n_samples))
    return _items_for_indices(units, indices, max_chars), warnings


def _build_text_sample_items(text: str, n_samples: int, max_chars: int) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Chunk plain text and select N representative items capped at max_chars."""
    chunks = split_text_into_chunks(text, max_tokens_per_chunk=MAX_TOKENS_PER_CHUNK)
    total = len(chunks)
    if total < 1:
        raise ValueError("document too small for sampling")

    warnings: List[Dict[str, Any]] = []
    indices = select_sample_indices(total, n_samples)
    if len(indices) < n_samples:
        warnings.append(_small_document_warning(total, len(indices), n_samples))
    return _items_for_indices(chunks, indices, max_chars), warnings


def _pick_random_unused_index(total: int, exclude: Set[int]) -> Optional[int]:
    """Pick a random interior index (1..total-2) not in `exclude`. None if all used."""
    if total < 1:
        return None
    if total < 3:
        candidates = [i for i in range(total) if i not in exclude]
    else:
        candidates = [i for i in range(1, total - 1) if i not in exclude]
    if not candidates:
        return None
    return random.choice(candidates)


def _instantiate_provider(column: Dict[str, Any]):
    """Build an LLMProvider from a column descriptor.

    Resolves `__USE_ENV__` placeholders to the corresponding env variable.
    """
    provider = (column.get("provider") or "ollama").lower()
    env_var = _provider_env_var(provider)
    endpoint = column.get("api_endpoint") or column.get("endpoint")
    if provider in {"openai", "nim"}:
        default_endpoint = (
            OPENAI_API_ENDPOINT if provider == "openai" else NIM_API_ENDPOINT
        )
        api_key, _source = resolve_api_key_for_endpoint(
            column.get("api_key"),
            env_var,
            endpoint=endpoint or default_endpoint,
            default_endpoint=default_endpoint,
        )
    else:
        api_key = _resolve_api_key(column.get("api_key"), env_var) if env_var else None

    kwargs: Dict[str, Any] = {
        "model": column.get("model"),
        "context_window": int(column.get("context_window") or OLLAMA_NUM_CTX),
    }
    if api_key or provider in {"openai", "nim"}:
        kwargs["api_key"] = api_key
    if endpoint:
        kwargs["api_endpoint"] = endpoint

    return create_llm_provider(provider, **kwargs)


def _compute_cost_usd(
    provider: str,
    model: str,
    prompt_tokens: int,
    completion_tokens: int,
    *,
    prompt_cache_hit_tokens: int = 0,
    prompt_cache_miss_tokens: int = 0,
) -> Optional[float]:
    """Best-effort USD cost for one LLM call. Returns None if pricing unknown."""
    pricing_tier = (
        get_deepseek_pricing_status().pricing_tier
        if str(provider).lower() == "deepseek"
        else None
    )
    pricing = get_default_pricing(provider, model, pricing_tier=pricing_tier)
    if not pricing:
        return None
    _input_cost, _output_cost, total_cost = calculate_usage_cost(
        pricing,
        prompt_tokens=prompt_tokens,
        completion_tokens=completion_tokens,
        prompt_cache_hit_tokens=prompt_cache_hit_tokens,
        prompt_cache_miss_tokens=prompt_cache_miss_tokens,
    )
    return round(total_cost, 6)


async def _execute_cell(
    *,
    sample_id: str,
    row: int,
    col: int,
    phase: str,
    prompt_pair,
    column: Dict[str, Any],
    ref_text: str,
    state: "SampleStateManager",
    socketio,
) -> Optional[str]:
    """Run one LLM call for a cell and emit its result.

    Shared by the translate and refine phases — they differ only in the prompt
    pair and the text the length ratio is measured against (`ref_text`).
    Returns the cleaned output on success, or None on error/empty/cancel. Emits
    exactly one WebSocket event (done or error), except when cancelled.
    """
    if state.is_cancelled(sample_id):
        return None

    started = time.perf_counter()
    provider = None
    try:
        provider = _instantiate_provider(column)
        response = await await_llm_call(
            provider.generate,
            prompt_pair.user,
            provider=provider,
            request_timeout=SAMPLE_REQUEST_TIMEOUT,
            deadline=SAMPLE_REQUEST_TIMEOUT,
            system_prompt=prompt_pair.system,
        )
        latency_ms = int((time.perf_counter() - started) * 1000)

        if response is None or not response.content:
            _emit_cell(
                socketio, state, sample_id, row, col, phase,
                status="error",
                output=None,
                metrics={"latency_ms": latency_ms},
                error="LLM returned an empty response",
            )
            return None

        # Strip <TRANSLATION>...</TRANSLATION> wrapper (and any <think> block)
        # like the main translation flow does. Fall back to raw content if the
        # tags are missing — same semantics as `was_fallback`.
        extracted = provider.extract_translation(response.content)
        used_fallback = response.was_fallback
        if extracted is None or not extracted.strip():
            extracted = response.content
            used_fallback = True
        guarded = guard_llm_output(
            clean_text_artifacts(extracted.strip()),
            phase=f"sample_{phase}",
        )
        output_text = guarded.text

        cost = _compute_cost_usd(
            column.get("provider", "ollama"),
            column.get("model", ""),
            response.prompt_tokens,
            response.completion_tokens,
            prompt_cache_hit_tokens=response.prompt_cache_hit_tokens,
            prompt_cache_miss_tokens=response.prompt_cache_miss_tokens,
        )
        src_len = max(1, len(ref_text))
        length_ratio = round(len(output_text) / src_len, 3)

        _emit_cell(
            socketio, state, sample_id, row, col, phase,
            status="done",
            output=output_text,
            metrics={
                "latency_ms": latency_ms,
                "prompt_tokens": response.prompt_tokens,
                "completion_tokens": response.completion_tokens,
                "cost_usd": cost,
                "length_ratio": length_ratio,
                "was_fallback": used_fallback,
                "was_truncated": response.was_truncated,
                "output_guard_issues": [issue.to_dict() for issue in guarded.issues],
                "output_guard_scores": guarded.scores,
                "output_guard_changed": guarded.changed,
            },
            error=None,
        )
        return output_text
    except Exception as exc:
        latency_ms = int((time.perf_counter() - started) * 1000)
        _emit_cell(
            socketio, state, sample_id, row, col, phase,
            status="error",
            output=None,
            metrics={"latency_ms": latency_ms},
            error=str(exc),
        )
        return None
    finally:
        if provider is not None:
            try:
                await provider.close()
            except Exception:
                pass


async def _run_cell_translate(
    *,
    sample_id: str,
    row: int,
    col: int,
    item: Dict[str, Any],
    column: Dict[str, Any],
    source_language: str,
    target_language: str,
    prompt_options: Dict[str, Any],
    glossary_block: str = "",
    state: "SampleStateManager",
    socketio,
) -> Optional[str]:
    """Run a single translate call. Returns the translated text, or None."""
    prompt_pair = generate_translation_prompt(
        main_content=item["source_text"],
        context_before=item.get("context_before", ""),
        context_after=item.get("context_after", ""),
        previous_translation_context="",
        source_language=source_language,
        target_language=target_language,
        has_placeholders=False,
        prompt_options=prompt_options,
        glossary_block=glossary_block,
    )
    return await _execute_cell(
        sample_id=sample_id, row=row, col=col, phase="translate",
        prompt_pair=prompt_pair, column=column, ref_text=item["source_text"],
        state=state, socketio=socketio,
    )


async def _run_cell_refine(
    *,
    sample_id: str,
    row: int,
    col: int,
    draft_text: str,
    item: Dict[str, Any],
    column: Dict[str, Any],
    target_language: str,
    prompt_options: Dict[str, Any],
    glossary_block: str = "",
    phase: str = "refine",
    state: "SampleStateManager",
    socketio,
) -> Optional[str]:
    """Run a single refine call. Emits one WebSocket event when done."""
    prompt_pair = generate_refinement_prompt(
        draft_translation=draft_text,
        context_before=item.get("context_before", ""),
        context_after=item.get("context_after", ""),
        previous_refined_context="",
        target_language=target_language,
        has_placeholders=False,
        prompt_options=prompt_options,
        glossary_block=glossary_block,
        # A preset's refinement section reaches the refine prompt via
        # `additional_instructions` (prompt_options alone isn't read for it).
        additional_instructions=(prompt_options.get("refinement_instructions") or ""),
    )
    return await _execute_cell(
        sample_id=sample_id, row=row, col=col, phase=phase,
        prompt_pair=prompt_pair, column=column, ref_text=draft_text,
        state=state, socketio=socketio,
    )


def _emit_cell(socketio, state, sample_id, row, col, phase, *, status, output, metrics, error):
    """Persist the cell result in state and emit it over WebSocket."""
    state.update_cell(
        sample_id, row, col, phase,
        status=status, output=output, metrics=metrics, error=error,
    )
    if socketio is None:
        return
    payload = {
        "sample_id": sample_id,
        "type": "cell_done" if status == "done" else "cell_error",
        "row": row,
        "col": col,
        "phase": phase,
        "output": output,
        "metrics": metrics or {},
        "error": error,
    }
    try:
        socketio.emit("sample_update", payload, namespace="/")
    except Exception as exc:
        logger.error("sample_update emit failed for %s: %s", sample_id, exc)


def _activate_column_profile(opts: Dict[str, Any], column: Dict[str, Any]) -> None:
    """Activate the selected book profile for a Sample/Quick Test column."""
    profile_id = str(column.get("profile_id") or opts.get("profile_id") or "").strip()
    if not profile_id:
        return

    opts["editorial_mode"] = "book_profile"
    opts["profile_id"] = profile_id
    opts["preserve_author_voice"] = True
    opts["use_profile_glossary"] = True
    opts["allow_common_glossary"] = True
    opts["allow_cross_profile_glossary"] = False
    opts["glossary_suggestions_enabled"] = True
    opts["auto_approve_glossary_suggestions"] = False
    opts["min_glossary_suggestion_confidence"] = 0.92
    opts["avoid_hardcoded_editorial_rules"] = True

    try:
        from src.core.book_profiles import load_book_profile

        profile = load_book_profile(profile_id, allow_missing=True)
    except Exception:
        profile = None
    if profile is None:
        return

    if profile.target_locale:
        opts.setdefault("target_locale", profile.target_locale)
    opts["preserve_author_voice"] = profile.preserve_author_voice
    opts["allow_common_glossary"] = profile.allow_common_glossary
    opts["allow_cross_profile_glossary"] = profile.allow_cross_profile_glossary
    opts.setdefault("modernization_strength", profile.modernization_strength or "high")
    opts.setdefault("min_dimension_score", profile.min_dimension_score)
    opts.setdefault("max_repair_rounds", profile.max_repair_rounds)


def _column_prompt_options(base_options: Dict[str, Any], column: Dict[str, Any]) -> Dict[str, Any]:
    """Merge the run-wide prompt_options with a column's own custom-instruction
    preset (a file in Custom_Instructions/). A per-column preset overrides the
    run-wide `custom_instructions` (translation phase) and `refinement_instructions`
    (refine phase). Best-effort: an unsafe/missing/empty file is ignored, so the
    column falls back to the run-wide options.
    """
    opts = dict(base_options or {})
    column_options = column.get("prompt_options")
    if isinstance(column_options, dict):
        opts.update(column_options)
    _activate_column_profile(opts, column)
    filename = (column.get("custom_instruction_file") or "").strip()
    if not filename:
        return opts
    if not is_safe_filename(filename):
        logger.warning("sample: ignoring unsafe custom instruction filename %r", filename)
        return opts
    try:
        ci_dir = Path(os.getcwd()) / "Custom_Instructions"
        loaded = load_custom_instructions(filename, ci_dir)
        translation = loaded.get("translation")
        refinement = loaded.get("refinement")
        if translation:
            opts["custom_instructions"] = translation
        if refinement:
            opts["refinement_instructions"] = refinement
    except Exception as exc:
        logger.warning("sample: failed to load custom instructions %r: %s", filename, exc)
    return opts


_TEXT_TRANSFORM_OPTION_KEYS = {
    "text_transform_mode",
    "text_transform_label",
    "text_transform_profile",
    "refinement_instructions",
    "transform_guard",
    "transform_auditor_model",
    "transform_repair_attempts",
    "transform_fallback",
}


def _translation_options_without_transform(options: Dict[str, Any]) -> Dict[str, Any]:
    """Keep profile/glossary translation context while preventing transform
    instructions from hijacking the first translation pass.
    """
    cleaned = dict(options or {})
    for key in _TEXT_TRANSFORM_OPTION_KEYS:
        cleaned.pop(key, None)
    return cleaned


def _load_column_glossary(glossary_id: Any) -> Optional[Dict[str, Any]]:
    """Load a glossary by id into {terms_dict, term_metadata, target_language},
    or None when there's no/invalid glossary. Best-effort: any failure (missing
    id, store error, empty glossary) yields None so the column runs without one.
    """
    if not glossary_id:
        return None
    try:
        from src.api.translation_state import get_glossary_store  # lazy: avoid cycle
        glossary = get_glossary_store().get_glossary(int(glossary_id))
    except Exception as exc:
        logger.warning("sample: failed to load glossary %r: %s", glossary_id, exc)
        return None
    if not glossary or not glossary.terms:
        return None
    return {
        "terms_dict": glossary.terms_dict,
        "term_metadata": {
            term.source_term: {"category": term.category or ""}
            for term in glossary.terms if term.source_term
        },
        "target_language": glossary.target_language or "",
    }


def _sample_glossary_purpose(mode: str, prompt_options: Dict[str, Any], phase: str) -> str:
    """Return the glossary matching/rendering purpose for a sample cell."""
    if (prompt_options or {}).get("text_transform_mode"):
        return "transformation"
    normalized_phase = str(phase or "").strip().lower()
    if normalized_phase in {"refine", "refinement"} or str(mode or "").strip().lower() == "refine":
        return "refinement"
    return "translation"


def _glossary_config_for_purpose(purpose: str) -> GlossaryConfig:
    if str(purpose or "").strip().lower() in {"refinement", "transformation"}:
        return GlossaryConfig(case_sensitive=False, accent_insensitive=True)
    return GlossaryConfig()


def _glossary_block_for(
    glossary_data: Optional[Dict[str, Any]],
    text: str,
    *,
    prompt_options: Optional[Dict[str, Any]] = None,
    purpose: str = "translation",
) -> str:
    """Build the per-cell glossary blocks for the active sample phase."""
    blocks: List[str] = []
    try:
        if glossary_data:
            filtered, _capped = filter_glossary_for_purpose(
                text,
                glossary_data["terms_dict"],
                _glossary_config_for_purpose(purpose),
                purpose,
            )
            if filtered:
                blocks.append(
                    build_glossary_block(
                        filtered_terms=filtered,
                        target_language=glossary_data["target_language"],
                        term_metadata=glossary_data["term_metadata"],
                        purpose=purpose,
                    ).strip()
                )

        profile_block = build_profile_glossary_block(
            text,
            prompt_options,
            purpose=purpose,
        )
        if profile_block:
            blocks.append(profile_block.strip())

        return "\n\n".join(block for block in blocks if block)
    except Exception as exc:
        logger.warning("sample: failed to build glossary block: %s", exc)
        return ""


async def _run_sample_async(
    *,
    sample_id: str,
    items: List[Dict[str, Any]],
    columns: List[Dict[str, Any]],
    mode: str,
    source_language: str,
    target_language: str,
    prompt_options: Dict[str, Any],
    state: "SampleStateManager",
    socketio,
    skip_cells: Optional[set] = None,
) -> None:
    """Run all N×K cells in parallel under a concurrency semaphore.

    `skip_cells` is a set of (row, col) tuples whose LLM calls must be skipped
    — used by the cross-run cache: when the client already has a cached result
    for that cell, we avoid spending tokens on it.
    """
    skip = skip_cells or set()
    sem = asyncio.Semaphore(min(SAMPLE_CONCURRENCY_CAP, max(1, len(items) * len(columns))))
    total_cells = len(items) * len(columns)
    logger.info(
        "sample %s: worker started mode=%s cells=%s skipped=%s timeout=%ss",
        sample_id,
        mode,
        total_cells,
        len(skip),
        SAMPLE_REQUEST_TIMEOUT,
    )

    # Resolve each column's custom-instruction preset and glossary once, so every
    # cell of that column reuses them (the glossary block is still filtered per
    # cell against that cell's source text).
    column_options = [_column_prompt_options(prompt_options, c) for c in columns]
    column_glossaries = [_load_column_glossary(c.get("glossary_id")) for c in columns]

    async def cell_task(row: int, col: int):
        async with sem:
            if state.is_cancelled(sample_id):
                return
            if (row, col) in skip:
                return
            item = items[row]
            column = columns[col]
            cell_options = column_options[col]
            if mode == "refine":
                purpose = _sample_glossary_purpose(mode, cell_options, "refine")
                glossary_block = _glossary_block_for(
                    column_glossaries[col],
                    item["source_text"],
                    prompt_options=cell_options,
                    purpose=purpose,
                )
                # Treat the source extract as the draft to refine.
                await _run_cell_refine(
                    sample_id=sample_id, row=row, col=col,
                    draft_text=item["source_text"], item=item, column=column,
                    target_language=target_language,
                    prompt_options=cell_options,
                    glossary_block=glossary_block,
                    state=state, socketio=socketio,
                )
                return

            translation_options = _translation_options_without_transform(cell_options)
            translation_purpose = _sample_glossary_purpose(mode, translation_options, "translate")
            translation_glossary_block = _glossary_block_for(
                column_glossaries[col],
                item["source_text"],
                prompt_options=translation_options,
                purpose=translation_purpose,
            )
            draft = await _run_cell_translate(
                sample_id=sample_id, row=row, col=col,
                item=item, column=column,
                source_language=source_language, target_language=target_language,
                prompt_options=translation_options,
                glossary_block=translation_glossary_block,
                state=state, socketio=socketio,
            )

            column_wants_refine = bool(column.get("refine_after_translate"))
            should_refine = (
                draft
                and not state.is_cancelled(sample_id)
                and (
                    column_wants_refine
                    or (mode == "translate_refine" and column.get("refine_after_translate") is None)
                )
            )
            if should_refine:
                refinement_purpose = _sample_glossary_purpose(mode, cell_options, "refine")
                refinement_glossary_text = "\n".join(
                    part for part in (item["source_text"], draft) if part
                )
                refinement_glossary_block = _glossary_block_for(
                    column_glossaries[col],
                    refinement_glossary_text,
                    prompt_options=cell_options,
                    purpose=refinement_purpose,
                )
                await _run_cell_refine(
                    sample_id=sample_id, row=row, col=col,
                    draft_text=draft, item=item, column=column,
                    target_language=target_language,
                    prompt_options=cell_options,
                    glossary_block=refinement_glossary_block,
                    state=state, socketio=socketio,
                )

    await asyncio.gather(
        *(cell_task(r, c) for r in range(len(items)) for c in range(len(columns))),
        return_exceptions=True,
    )

    final_status = "stopped" if state.is_cancelled(sample_id) else "completed"
    state.set_status(sample_id, final_status)
    logger.info("sample %s: worker finished status=%s", sample_id, final_status)
    if socketio is not None:
        try:
            socketio.emit(
                "sample_update",
                {
                    "sample_id": sample_id,
                    "type": "sample_stopped" if final_status == "stopped" else "sample_done",
                },
                namespace="/",
            )
        except Exception as exc:
            logger.error("sample_update final emit failed for %s: %s", sample_id, exc)


def _spawn_sample_thread(coro_factory):
    """Run an async coroutine in a fresh thread with its own event loop."""
    def runner():
        loop = asyncio.new_event_loop()
        try:
            loop.run_until_complete(coro_factory())
        finally:
            # Finalize any async generators still open (e.g. the httpx streaming
            # response behind `async for line in response.aiter_lines()` in the
            # LLM providers) before closing the loop. Without this, their cleanup
            # tasks are destroyed mid-flight and asyncio logs
            # "Task was destroyed but it is pending!". Mirrors glossary_routes.
            try:
                loop.run_until_complete(loop.shutdown_asyncgens())
            except Exception:
                pass
            loop.close()
    thread = threading.Thread(target=runner, daemon=True)
    thread.start()
    return thread


def create_sample_blueprint(sample_state_manager, socketio=None, output_dir=None):
    """Create the sample blueprint.

    Args:
        sample_state_manager: Instance of SampleStateManager.
        socketio: SocketIO instance, used to emit `sample_update` events.
    """
    bp = Blueprint("sample", __name__)
    managed_uploads = Path(output_dir or OUTPUT_DIR) / "uploads"

    def _validate_file(data: Dict[str, Any]) -> Tuple[Optional[str], Optional[str], Optional[Tuple[Any, int]]]:
        """Common file_path + file_type validation. Returns (path, type, err)."""
        file_path = data.get("file_path")
        if not file_path:
            return None, None, (jsonify({"error": "Missing field: file_path"}), 400)
        try:
            managed_file = PathValidator.resolve_managed_file(file_path, [managed_uploads])
        except ValueError:
            return None, None, (jsonify({"error": "File path is outside managed uploads"}), 403)
        except FileNotFoundError:
            return None, None, (jsonify({"error": "Uploaded file not found"}), 404)
        file_path = str(managed_file)
        try:
            detected = detect_file_type(file_path)
        except Exception as exc:
            return None, None, (jsonify({"error": f"Cannot detect file type: {exc}"}), 400)
        file_type = (data.get("file_type") or detected).lower()
        if file_type != detected:
            return None, None, (jsonify({
                "error": f"File type mismatch: client said {file_type!r}, server detected {detected!r}",
            }), 400)
        return file_path, file_type, None

    @bp.route("/api/sample/initialize", methods=["POST"])
    def initialize_samples():
        """Sample N initial extracts from a freshly uploaded file.

        Called by the client right after upload so the user can preview the
        selected blocks before spending any LLM tokens. Returns items with the
        same shape /api/sample/run produces, but without creating a sample_id
        and without spawning any background work.
        """
        data = request.get_json(silent=True) or {}
        file_path, file_type, err = _validate_file(data)
        if err is not None:
            return err

        try:
            n_samples = _clamp_int(data.get("n_samples"), DEFAULT_N_SAMPLES, 2, 20)
            max_chars = _clamp_int(data.get("max_chars"), DEFAULT_MAX_CHARS, 50, 2000)
        except (TypeError, ValueError):
            return jsonify({"error": "n_samples and max_chars must be integers"}), 400

        try:
            units = _load_source_units(file_path, file_type)
            total = len(units)
            if total < 1:
                return jsonify({"error": "document too small for sampling"}), 400
            indices = select_sample_indices(total, n_samples)
            items = _items_for_indices(units, indices, max_chars)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except Exception as exc:
            return jsonify({"error": f"Failed to initialize samples: {exc}"}), 500

        warnings: List[Dict[str, Any]] = []
        if len(indices) < n_samples:
            warnings.append(_small_document_warning(total, len(indices), n_samples))

        public_items = [
            {"index": it["index"], "source_text": it["source_text"], "truncated": it["truncated"]}
            for it in items
        ]
        return jsonify({"items": public_items, "total": total, "warnings": warnings})

    @bp.route("/api/sample/extract", methods=["POST"])
    def extract_random_sample():
        """Pick a random extract not in `exclude_indices` (server-side RNG).

        Used by the "Add a sample" button to grow the user's curated sample
        list. Returns 409 when the document has no remaining interior index.
        """
        data = request.get_json(silent=True) or {}
        file_path, file_type, err = _validate_file(data)
        if err is not None:
            return err

        try:
            max_chars = _clamp_int(data.get("max_chars"), DEFAULT_MAX_CHARS, 50, 2000)
        except (TypeError, ValueError):
            return jsonify({"error": "max_chars must be an integer"}), 400

        raw_excl = data.get("exclude_indices") or []
        exclude: Set[int] = set()
        for v in raw_excl:
            try:
                exclude.add(int(v))
            except (TypeError, ValueError):
                continue

        try:
            units = _load_source_units(file_path, file_type)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except Exception as exc:
            return jsonify({"error": f"Failed to load source: {exc}"}), 500

        total = len(units)
        if total < 1:
            return jsonify({"error": "document too small for sampling"}), 400

        idx = _pick_random_unused_index(total, exclude)
        if idx is None:
            return jsonify({"error": "no_more_indices", "total": total}), 409

        items = _items_for_indices(units, [idx], max_chars)
        if not items:
            return jsonify({"error": "failed to build item"}), 500
        it = items[0]
        return jsonify({
            "item": {
                "index": it["index"],
                "source_text": it["source_text"],
                "truncated": it["truncated"],
            },
            "total": total,
        })

    @bp.route("/api/sample/run", methods=["POST"])
    def start_sample_run():
        data = request.get_json(silent=True) or {}

        # Required fields. `source_language` may be empty: we auto-detect it
        # from the uploaded file's content (mirrors the Translate-tab behavior).
        for field in ("file_path", "file_type", "target_language", "columns"):
            if field not in data or data[field] in (None, "", []):
                return jsonify({"error": f"Missing or empty field: {field}"}), 400

        # File existence + type detection/mismatch (shared with initialize/extract).
        file_path, file_type, err = _validate_file(data)
        if err is not None:
            return err

        mode = (data.get("mode") or "translate").lower()
        if mode not in ("translate", "refine", "translate_refine"):
            return jsonify({"error": f"Invalid mode: {mode}"}), 400

        try:
            n_samples = _clamp_int(data.get("n_samples"), DEFAULT_N_SAMPLES, 2, 20)
            max_chars = _clamp_int(data.get("max_chars"), DEFAULT_MAX_CHARS, 50, 2000)
        except (TypeError, ValueError):
            return jsonify({"error": "n_samples and max_chars must be integers"}), 400

        columns_raw = data["columns"]
        if not isinstance(columns_raw, list) or len(columns_raw) < 1:
            return jsonify({"error": "columns must be a non-empty list"}), 400

        # Build sample items.
        #
        # Two paths:
        #  - `items` provided by the client → user already curated the sample
        #    set (initialize + add/remove). We honor the indices and the
        #    client-supplied source_text, but re-derive context_before/after
        #    server-side (deterministic given the file).
        #  - `items` missing → fall back to the legacy auto-sampling path.
        warnings: List[Dict[str, Any]] = []
        client_items = data.get("items")
        try:
            if client_items is not None:
                if not isinstance(client_items, list) or not client_items:
                    return jsonify({"error": "items must be a non-empty list"}), 400
                units = _load_source_units(file_path, file_type)
                items = []
                total = len(units)
                for raw in client_items:
                    if not isinstance(raw, dict):
                        continue
                    try:
                        idx = int(raw.get("index"))
                    except (TypeError, ValueError):
                        continue
                    if idx < 0 or idx >= total:
                        continue
                    source_text = raw.get("source_text")
                    if not isinstance(source_text, str) or not source_text.strip():
                        continue
                    items.append({
                        "index": idx,
                        "source_text": source_text,
                        "truncated": bool(raw.get("truncated")),
                        "context_before": units[idx].get("context_before", ""),
                        "context_after": units[idx].get("context_after", ""),
                    })
                if not items:
                    return jsonify({"error": "no valid items provided"}), 400
            elif file_type == "srt":
                items, warnings = _build_srt_sample_blocks(file_path, n_samples, max_chars)
            else:
                text = _extract_plain_text(file_path, file_type)
                if not text or not text.strip():
                    return jsonify({"error": "File is empty or unreadable"}), 400
                items, warnings = _build_text_sample_items(text, n_samples, max_chars)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except Exception as exc:
            return jsonify({"error": f"Failed to prepare samples: {exc}"}), 500

        # Normalize columns and create state entry
        columns = []
        for raw in columns_raw:
            prompt_options_raw = raw.get("prompt_options")
            if not isinstance(prompt_options_raw, dict):
                prompt_options_raw = {}
            refine_after_translate = (
                bool(raw.get("refine_after_translate"))
                if "refine_after_translate" in raw
                else None
            )
            columns.append({
                "provider": (raw.get("provider") or "ollama").lower(),
                "model": raw.get("model") or "",
                "api_key": raw.get("api_key"),
                "api_endpoint": raw.get("api_endpoint") or raw.get("endpoint"),
                "custom_instruction_file": raw.get("custom_instruction_file") or "",
                "glossary_id": raw.get("glossary_id") or None,
                "profile_id": raw.get("profile_id") or None,
                "variant_key": raw.get("variant_key") or "",
                "label": raw.get("label") or "",
                "variant_summary": raw.get("variant_summary") or "",
                "refine_after_translate": refine_after_translate,
                "prompt_options": prompt_options_raw,
            })

        sample_id = f"sample_{int(time.time() * 1000)}_{uuid.uuid4().hex[:6]}"
        sample_state_manager.create(sample_id, items, columns, mode)
        logger.info(
            "sample %s: prepared mode=%s items=%s columns=%s file_type=%s",
            sample_id,
            mode,
            len(items),
            len(columns),
            file_type,
        )

        # Items exposed to the client must not leak context_before/context_after
        # — those are kept server-side only and used to enrich prompts.
        public_items = [
            {"index": it["index"], "source_text": it["source_text"], "truncated": it["truncated"]}
            for it in items
        ]
        public_columns = [
            {k: v for k, v in col.items() if k != "api_key"}
            for col in columns
        ]

        source_language = (data.get("source_language") or "").strip()
        target_language = data["target_language"]
        # Auto-detect source language from file content when the user leaves
        # the picker on "Auto-detect" (same UX as the Translate tab).
        if not source_language:
            try:
                with open(file_path, "rb") as fh:
                    file_bytes = fh.read()
                detected_name, confidence = LanguageDetector.detect_language_from_file(
                    file_bytes, os.path.basename(file_path)
                )
                if detected_name:
                    source_language = detected_name
                    warnings.append({
                        "code": "warning_lang_autodetected",
                        "params": {"lang": detected_name, "confidence": round(confidence * 100)},
                    })
            except Exception as exc:
                logger.warning("sample: language auto-detection failed: %s", exc)
        if not source_language:
            return jsonify({
                "error": "Could not auto-detect source language; please pick one manually.",
            }), 400

        if mode == "refine":
            target_language = source_language
        prompt_options = data.get("prompt_options") or {}

        # `defer_dispatch=true` lets the client read the items first, compute
        # which cells are already cached, then call /dispatch with skip_cells.
        defer_dispatch = bool(data.get("defer_dispatch"))
        if defer_dispatch:
            sample_state_manager.set_status(sample_id, "prepared")

        # Stash the run parameters so /dispatch can pick them up. Pending state
        # entries are already created by sample_state_manager.create().
        sample_state_manager.set_run_context(sample_id, {
            "items": items,
            "columns": columns,
            "mode": mode,
            "source_language": source_language,
            "target_language": target_language,
            "prompt_options": prompt_options,
        })

        if not defer_dispatch:
            async def _runner():
                await _run_sample_async(
                    sample_id=sample_id,
                    items=items,
                    columns=columns,
                    mode=mode,
                    source_language=source_language,
                    target_language=target_language,
                    prompt_options=prompt_options,
                    state=sample_state_manager,
                    socketio=socketio,
                )

            _spawn_sample_thread(_runner)

        return jsonify({
            "sample_id": sample_id,
            "items": public_items,
            "columns": public_columns,
            "mode": mode,
            "warnings": warnings,
            "deferred": defer_dispatch,
        })

    @bp.route("/api/sample/<sample_id>/dispatch", methods=["POST"])
    def dispatch_sample_run(sample_id):
        """Start the LLM work for a previously prepared (deferred) run.

        Body: { skip_cells: [[row, col], ...] }. Cells in skip_cells are not
        sent to the LLM — the client already has them cached from an earlier
        run with identical parameters.
        """
        if not sample_state_manager.exists(sample_id):
            return jsonify({"error": "Sample run not found"}), 404

        run_ctx = sample_state_manager.get_run_context(sample_id)
        if run_ctx is None:
            return jsonify({"error": "Sample run has no pending dispatch context"}), 409

        payload = request.get_json(silent=True) or {}
        raw_skip = payload.get("skip_cells") or []
        skip: set = set()
        for pair in raw_skip:
            if isinstance(pair, (list, tuple)) and len(pair) == 2:
                try:
                    skip.add((int(pair[0]), int(pair[1])))
                except (TypeError, ValueError):
                    continue

        async def _runner():
            await _run_sample_async(
                sample_id=sample_id,
                items=run_ctx["items"],
                columns=run_ctx["columns"],
                mode=run_ctx["mode"],
                source_language=run_ctx["source_language"],
                target_language=run_ctx["target_language"],
                prompt_options=run_ctx["prompt_options"],
                state=sample_state_manager,
                socketio=socketio,
                skip_cells=skip,
            )

        sample_state_manager.set_status(sample_id, "running")
        logger.info("sample %s: dispatch started skipped=%s", sample_id, len(skip))
        _spawn_sample_thread(_runner)
        return jsonify({"message": "Dispatch started", "skipped": len(skip)}), 200

    @bp.route("/api/sample/runs", methods=["GET"])
    def list_sample_runs():
        """List currently retained Sample & Compare runs for diagnostics."""
        return jsonify({"runs": sample_state_manager.list_summaries()})

    @bp.route("/api/sample/<sample_id>/stop", methods=["POST"])
    def stop_sample_run(sample_id):
        if not sample_state_manager.exists(sample_id):
            return jsonify({"error": "Sample run not found"}), 404
        sample_state_manager.cancel(sample_id)
        return jsonify({"message": "Sample run stopped"}), 200

    @bp.route("/api/sample/<sample_id>", methods=["GET"])
    def get_sample_run(sample_id):
        snapshot = sample_state_manager.get(sample_id)
        if snapshot is None:
            return jsonify({"error": "Sample run not found"}), 404
        # Strip API keys defensively before returning
        snapshot["columns"] = [
            {k: v for k, v in col.items() if k != "api_key"}
            for col in snapshot.get("columns", [])
        ]
        # Strip server-only context fields from items
        snapshot["items"] = [
            {"index": it["index"], "source_text": it["source_text"], "truncated": it["truncated"]}
            for it in snapshot.get("items", [])
        ]
        run_context = snapshot.get("run_context") or {}
        snapshot["run_context"] = {
            "mode": run_context.get("mode"),
            "source_language": run_context.get("source_language"),
            "target_language": run_context.get("target_language"),
            "prompt_options": run_context.get("prompt_options") or {},
        }
        return jsonify(snapshot)

    return bp
