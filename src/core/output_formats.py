"""Final output format helpers.

Translation/refinement engines write their native format first. This module
optionally converts that result into a user-requested delivery format.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import html
import posixpath
from pathlib import Path
import re
import uuid
import zipfile
from urllib.parse import unquote

from lxml import etree

from src.core.editorial_assembly import EditorialAssemblyAgent
from src.core.layout_sanitizer import sanitize_extracted_text, strip_inline_style_markers
from src.utils.archive_safety import validate_zip_archive


SUPPORTED_OUTPUT_FORMATS = ("auto", "txt", "docx", "pdf", "epub")

_FORMAT_EXTENSIONS = {
    "txt": ".txt",
    "docx": ".docx",
    "pdf": ".pdf",
    "epub": ".epub",
}


def native_output_format(file_type: str) -> str:
    """Return the pipeline's native output format for an input file type."""
    ft = (file_type or "txt").lower()
    if ft == "pdf":
        return "txt"
    if ft in {"txt", "srt", "epub", "docx"}:
        return ft
    return "txt"


def normalize_output_format(value: str | None) -> str:
    """Normalize a requested output format; unknown values become ``auto``."""
    fmt = (value or "auto").strip().lower().lstrip(".")
    return fmt if fmt in SUPPORTED_OUTPUT_FORMATS else "auto"


def format_extension(output_format: str, fallback: str = ".txt") -> str:
    return _FORMAT_EXTENSIONS.get(output_format, fallback)


def ensure_output_extension(filename: str, output_format: str) -> str:
    """Force the canonical extension for a resolved output format."""
    fmt = normalize_output_format(output_format)
    if fmt == "auto":
        return filename
    desired_ext = format_extension(fmt)
    path = Path(filename)
    stem = path.name[:-len(path.suffix)] if path.suffix else path.name
    return f"{stem}{desired_ext}"


def requested_format_for_job(file_type: str, output_format: str | None) -> str:
    """Resolve ``auto`` to the native pipeline output format."""
    fmt = normalize_output_format(output_format)
    return native_output_format(file_type) if fmt == "auto" else fmt


def convert_output_file(
    source_path: str | Path,
    destination_path: str | Path,
    output_format: str,
    structure_hints: dict | None = None,
) -> None:
    """Convert a completed native output file to ``output_format``."""
    fmt = normalize_output_format(output_format)
    if fmt == "auto":
        fmt = _format_from_suffix(Path(destination_path).suffix)
    if fmt not in {"txt", "docx", "pdf", "epub"}:
        raise ValueError(f"Unsupported output format: {output_format}")

    source = Path(source_path)
    destination = Path(destination_path)
    text = extract_readable_text(source)

    if fmt == "txt":
        from src.utils.text_encoding import encode_utf8_text_download

        destination.write_bytes(encode_utf8_text_download(text))
    elif fmt == "docx":
        _write_docx(text, destination)
    elif fmt == "pdf":
        _write_pdf(text, destination)
    elif fmt == "epub":
        _write_epub(text, destination, title=destination.stem, structure_hints=structure_hints)


def write_text_as_output(
    text: str,
    destination_path: str | Path,
    output_format: str,
    structure_hints: dict | None = None,
) -> None:
    """Write already-prepared text as one of the supported delivery formats."""
    fmt = normalize_output_format(output_format)
    if fmt == "auto":
        fmt = _format_from_suffix(Path(destination_path).suffix)
    if fmt not in {"txt", "docx", "pdf", "epub"}:
        raise ValueError(f"Unsupported output format: {output_format}")

    destination = Path(destination_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "txt":
        from src.utils.text_encoding import encode_utf8_text_download

        destination.write_bytes(encode_utf8_text_download(text or ""))
    elif fmt == "docx":
        _write_docx(text or "", destination)
    elif fmt == "pdf":
        _write_pdf(text or "", destination)
    elif fmt == "epub":
        _write_epub(text or "", destination, title=destination.stem, structure_hints=structure_hints)


def extract_readable_text(path: str | Path) -> str:
    """Extract readable text from a translated output file."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    try:
        from src.utils.file_detector import detect_file_type

        detected_type = detect_file_type(str(file_path))
    except Exception:
        detected_type = None

    if (
        detected_type in {"txt", "srt"}
        or suffix in {".txt", ".text", ".md", ".markdown", ".log", ".csv", ".json", ".xml", ".html", ".htm", ".yaml", ".yml", ".srt"}
    ):
        return file_path.read_text(encoding="utf-8", errors="replace").strip()

    if detected_type == "pdf" or suffix == ".pdf":
        from src.core.pdf import extract_pdf_text
        return extract_pdf_text(file_path).strip()

    if detected_type == "docx" or suffix == ".docx":
        return _extract_docx_readable_text(file_path).strip()

    if detected_type == "epub" or suffix == ".epub":
        return _extract_epub_text(file_path).strip()

    return file_path.read_text(encoding="utf-8", errors="replace").strip()


def _extract_docx_readable_text(file_path: Path) -> str:
    """Extract DOCX body content as sanitized editorial text.

    This intentionally ignores Word headers, footers, margins, pagination, and
    run-level styling. Tables are kept as readable cell blocks so their content
    is not lost, but the LLM receives content rather than layout.
    """
    from src.core.docx.plain_extractor import extract_plain_paragraphs

    content = extract_plain_paragraphs(str(file_path))
    blocks = []
    table_cells = content.table_cell_indices()
    for index, value in enumerate(content.paragraphs_text):
        text, _style_count = strip_inline_style_markers(value or "")
        text = text.strip()
        if not text:
            continue
        if index in table_cells:
            blocks.append(text)
        else:
            blocks.append(text)

    cleaned, _report = sanitize_extracted_text(
        "\n\n".join(blocks),
        source_type="docx",
        strip_inline_style=True,
    )
    return cleaned


def build_epub_structure_candidates(
    text: str,
    document_title: str | None = None,
    *,
    max_preview_chars: int = 260,
) -> list[dict]:
    """Return compact EPUB section candidates for optional LLM adjudication."""
    candidates = []
    for index, section in enumerate(_build_output_sections(text, document_title), start=1):
        preview = ""
        for paragraph in section.paragraphs:
            if paragraph.strip():
                preview = paragraph.strip()
                break
        candidates.append({
            "index": index,
            "title": section.title,
            "paragraph_count": len(section.paragraphs),
            "preview": preview[:max_preview_chars],
        })
    return candidates


def extract_epub_toc_titles(path: str | Path, *, limit: int = 160) -> list[str]:
    """Extract navigation titles from an EPUB's nav.xhtml or toc.ncx."""
    file_path = Path(path)
    if file_path.suffix.lower() != ".epub" or not file_path.exists():
        return []

    titles: list[str] = []
    seen: set[str] = set()

    def add(value: str | None) -> None:
        cleaned = _clean_heading(value or "")
        if not cleaned or cleaned in seen:
            return
        seen.add(cleaned)
        titles.append(cleaned)

    with zipfile.ZipFile(file_path, "r") as zf:
        validate_zip_archive(zf)
        for name in zf.namelist():
            lower = name.lower()
            if lower.endswith("toc.ncx"):
                try:
                    root = etree.fromstring(zf.read(name), etree.XMLParser(recover=True, huge_tree=True))
                except Exception:
                    continue
                for node in root.xpath("//*[local-name()='navPoint']"):
                    labels = node.xpath(".//*[local-name()='navLabel']//*[local-name()='text']/text()")
                    if labels:
                        add(labels[0])
                    if len(titles) >= limit:
                        return titles

        for name in zf.namelist():
            lower = name.lower()
            if not lower.endswith((".xhtml", ".html", ".htm")):
                continue
            if "nav" not in lower and "toc" not in lower:
                continue
            try:
                root = etree.fromstring(zf.read(name), etree.XMLParser(recover=True, huge_tree=True))
            except Exception:
                continue
            for node in root.xpath("//*[local-name()='nav']//*[local-name()='a' or local-name()='span']"):
                add(" ".join(node.itertext()))
                if len(titles) >= limit:
                    return titles
    return titles


def _format_from_suffix(suffix: str) -> str:
    ext = (suffix or "").lower().lstrip(".")
    if ext in {"docx", "pdf", "epub"}:
        return ext
    return "txt"


def _extract_epub_text(file_path: Path) -> str:
    ordered_names: list[str] = []
    seen: set[str] = set()

    def add_name(name: str) -> None:
        normalized = posixpath.normpath(unquote(name)).lstrip("/")
        if (
            normalized
            and not normalized.startswith("../")
            and normalized in available
            and normalized not in seen
        ):
            ordered_names.append(normalized)
            seen.add(normalized)

    parts: list[str] = []
    with zipfile.ZipFile(file_path, "r") as zf:
        validate_zip_archive(zf)
        available = set(zf.namelist())
        opf_path = _find_epub_opf_path(zf)
        if opf_path:
            for name in _epub_spine_xhtml_names(zf, opf_path):
                add_name(name)

        if not ordered_names:
            # Fallback/coverage: append XHTML files only when no spine could be
            # resolved. If a valid spine exists, adding unspined nav/toc files
            # pollutes the translated book with duplicated index text.
            for name in zf.namelist():
                lower = name.lower()
                if not lower.endswith((".xhtml", ".html", ".htm")):
                    continue
                if _is_navigation_xhtml_name(lower):
                    continue
                add_name(name)

        for name in ordered_names:
            text = _extract_xhtml_body_text(zf.read(name))
            if text:
                parts.append(text)
    cleaned, _report = sanitize_extracted_text(
        "\n\n".join(parts),
        source_type="epub",
        strip_inline_style=True,
    )
    return cleaned


def _is_navigation_xhtml_name(lower_name: str) -> bool:
    base = posixpath.basename(lower_name)
    return (
        base in {"nav.xhtml", "nav.html", "toc.xhtml", "toc.html", "contents.xhtml", "contents.html"}
        or "/nav." in lower_name
        or "/toc." in lower_name
    )


def _find_epub_opf_path(zf: zipfile.ZipFile) -> str | None:
    """Return the package OPF path from META-INF/container.xml when present."""
    try:
        container = zf.read("META-INF/container.xml")
    except KeyError:
        container = b""

    if container:
        try:
            root = etree.fromstring(container)
            rootfiles = root.xpath("//*[local-name()='rootfile']")
            preferred = None
            for rootfile in rootfiles:
                full_path = rootfile.get("full-path")
                if not full_path:
                    continue
                if rootfile.get("media-type") == "application/oebps-package+xml":
                    preferred = full_path
                    break
                if preferred is None:
                    preferred = full_path
            if preferred:
                normalized = posixpath.normpath(unquote(preferred)).lstrip("/")
                if normalized in zf.namelist():
                    return normalized
        except Exception:
            pass

    for name in zf.namelist():
        if name.lower().endswith(".opf"):
            return name
    return None


def _epub_spine_xhtml_names(zf: zipfile.ZipFile, opf_path: str) -> list[str]:
    """Resolve XHTML/HTML files in EPUB spine order."""
    try:
        root = etree.fromstring(zf.read(opf_path))
    except Exception:
        return []

    manifest = {}
    for item in root.xpath("//*[local-name()='manifest']/*[local-name()='item']"):
        item_id = item.get("id")
        href = item.get("href")
        if item_id and href:
            manifest[item_id] = item

    opf_dir = posixpath.dirname(opf_path)
    ordered: list[str] = []
    for itemref in root.xpath("//*[local-name()='spine']/*[local-name()='itemref']"):
        item = manifest.get(itemref.get("idref") or "")
        if item is None:
            continue
        media_type = (item.get("media-type") or "").lower()
        href = item.get("href") or ""
        if media_type not in {"application/xhtml+xml", "text/html"}:
            continue
        ordered.append(posixpath.normpath(posixpath.join(opf_dir, unquote(href))))
    return ordered


def _extract_xhtml_body_text(data: bytes) -> str:
    try:
        parser = etree.XMLParser(
            encoding="utf-8",
            recover=True,
            remove_blank_text=False,
            huge_tree=True,
        )
        tree = etree.fromstring(data, parser)
    except Exception:
        try:
            tree = etree.fromstring(data, etree.HTMLParser())
        except Exception:
            return ""

    if tree is None:
        return ""

    body = _find_readable_html_body(tree)
    if body is None:
        return ""

    try:
        from src.core.epub.plain_extractor import extract_plain_paragraphs

        paragraphs, _tags, _images, _tables = extract_plain_paragraphs(body)
        text = "\n\n".join(p.strip() for p in paragraphs if p and p.strip())
    except Exception:
        text = ""

    if not text:
        text = etree.tostring(body, method="text", encoding="unicode")
        text = re.sub(r"\s+", " ", text).strip()
    return text


def _find_readable_html_body(tree: etree._Element) -> etree._Element | None:
    matches = tree.xpath(
        "//*[translate(local-name(), 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz')='body']"
    )
    if matches:
        return matches[0]
    local_name = etree.QName(tree).localname.lower() if isinstance(tree.tag, str) else ""
    if local_name in {"html", "section", "article", "div"}:
        return tree
    return None


@dataclass
class _OutputSection:
    title: str
    paragraphs: list[str]


_UPPER_RE_CHARS = "A-Z\\u00C0-\\u00D6\\u00D8-\\u00DE"
_LETTER_RE_CHARS = "A-Za-z\\u00C0-\\u024F"
_MONTH_RE = (
    "enero|febrero|marzo|abril|mayo|junio|julio|agosto|"
    "septiembre|setiembre|octubre|noviembre|diciembre|"
    "january|february|march|april|may|june|july|august|"
    "september|october|november|december"
)
_DATE_MARKER_RE = (
    rf"(?:c\.\s*)?\d{{3,4}}|"
    rf"(?:\d{{1,2}}\s+de\s+)?(?:{_MONTH_RE})\s+de\s+\d{{3,4}}|"
    rf"(?:{_MONTH_RE})\s+\d{{1,2}},?\s+\d{{3,4}}"
)
_DATE_TITLE_DELIMITER_RE = rf"(?:,\s*|:\s*[^.!?\n]{{0,70}},\s*)"
_LEADING_DATE_TITLE_RE = re.compile(
    rf"^(?P<title>[{_UPPER_RE_CHARS}][^.!?\n]{{6,180}}?{_DATE_TITLE_DELIMITER_RE}(?:{_DATE_MARKER_RE})\b)"
    rf"[.!?:;,-]*\s*(?P<rest>.*)$",
    re.IGNORECASE,
)
_EMBEDDED_DATE_TITLE_START_RE = re.compile(
    rf"([.!?\u2026][\"'\)\]\u00BB\u201D]?\s+)"
    rf"(?=[{_UPPER_RE_CHARS}][^.!?\n]{{6,180}}?{_DATE_TITLE_DELIMITER_RE}(?:{_DATE_MARKER_RE})\b)",
    re.IGNORECASE,
)
_AUTHOR_WORD_RE = rf"(?:[{_UPPER_RE_CHARS}][{_LETTER_RE_CHARS}'.\u2019-]*|[{_UPPER_RE_CHARS}]\.)"
_EMBEDDED_AUTHOR_DATE_TITLE_START_RE = re.compile(
    rf"(\b(?!(?:El|La|Los|Las|Un|Una|The|A|An)\s+)"
    rf"{_AUTHOR_WORD_RE}(?:\s+{_AUTHOR_WORD_RE}){{1,4}}\s+)"
    rf"(?=[{_UPPER_RE_CHARS}][^.!?\n]{{4,160}}?{_DATE_TITLE_DELIMITER_RE}(?:{_DATE_MARKER_RE})\b)",
)
_CONTENTS_HEADING_RE = re.compile(
    r"^(?:contents|contenido|tabla de contenidos?|indice|\u00EDndice)\b",
    re.IGNORECASE,
)
_DECLARED_HEADING_RE = re.compile(
    r"^(?:chapter|capitulo|cap\u00EDtulo|part|parte|book|libro|section|secci\u00F3n)"
    r"(?:\s+[\wIVXLCDMivxlcdm.\-]+)?(?:\s*[:.\-]\s*.+|\s+.+)?$",
    re.IGNORECASE,
)
_FRONT_BACK_HEADING_RE = re.compile(
    r"^(?:prologo|pr\u00F3logo|prefacio|introducci\u00F3n|introduccion|"
    r"epilogo|ep\u00EDlogo|apendice|ap\u00E9ndice|bibliograf\u00EDa|"
    r"bibliografia|notas|agradecimientos|glosario|contents|contenido|"
    r"(?:indice|\u00EDndice)(?:\s+de\s+[\w\u00C0-\u024F ]{2,60})?)$",
    re.IGNORECASE,
)
_CHRONICLE_MARKER_RE = re.compile(
    r"\s+(?=(?:"
    r"El\s+(?:d[i\u00ED]a|primero|segundo|tercero|cuarto|quinto|sexto|"
    r"s[e\u00E9]ptimo|octavo|noveno|d[e\u00E9]cimo|und[e\u00E9]cimo|"
    r"duod[e\u00E9]cimo|decimotercero|decimocuarto|decimoquinto|"
    r"diecis[e\u00E9]is|diecisiete|dieciocho|diecinueve|veinte|"
    r"veintiuno|veintid[o\u00F3]s|veintitr[e\u00E9]s|veinticuatro|"
    r"veinticinco|veintis[e\u00E9]is|veintisiete|veintiocho|"
    r"veintinueve|treinta|treinta y uno)\b|"
    r"Al\s+d[i\u00ED]a\s+siguiente\b|"
    r"A\s+la\s+ma\u00F1ana\s+siguiente\b"
    r"))",
    re.IGNORECASE,
)
_SENTENCE_BOUNDARY_RE = re.compile(
    rf"(?<=[.!?\u2026])\s+(?=[\(\"'\u00AB\u201C\u00BF\u00A1]?[{_UPPER_RE_CHARS}0-9])"
)


def _normalize_output_text(text: str) -> str:
    value = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    value = "".join(
        ch if ch in {"\n", "\t"} or ord(ch) >= 32 else " "
        for ch in value
    )
    # Common OCR split in years/page-like numbers: "162. 0" -> "1620".
    value = re.sub(r"(\b\d{3})[.,]\s+(\d\b)", r"\1\2", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"[ \t]*\n[ \t]*", "\n", value)
    value = re.sub(r"\n{4,}", "\n\n\n", value)
    return value.strip()


def _clean_heading(value: str) -> str:
    heading = re.sub(r"\s+", " ", value or "").strip()
    heading = heading.strip(" \t\r\n-:;,.!?")
    return heading or "Seccion"


def _clean_paragraph(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _raw_blocks(text: str) -> list[str]:
    normalized = _normalize_output_text(text)
    blocks = [b.strip() for b in re.split(r"\n\s*\n+", normalized) if b.strip()]
    if blocks:
        return blocks
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    if lines:
        return [" ".join(lines)]
    return []


def _split_embedded_section_starts(block: str) -> list[str]:
    """Split anthology entries that were glued to the prior paragraph."""
    if len(block) < 80:
        return [block]

    def starts_with_uppercase(index: int) -> bool:
        return bool(re.match(rf"\s*[{_UPPER_RE_CHARS}]", block[index:]))

    split_points = sorted({
        match.end(1)
        for pattern in (_EMBEDDED_DATE_TITLE_START_RE, _EMBEDDED_AUTHOR_DATE_TITLE_START_RE)
        for match in pattern.finditer(block)
        if match.end(1) > 0 and starts_with_uppercase(match.end(1))
    })
    if not split_points:
        return [block]

    parts: list[str] = []
    last = 0
    for split_at in split_points:
        if split_at <= last:
            continue
        previous = block[last:split_at].strip()
        if previous:
            parts.append(previous)
        last = split_at

    tail = block[last:].strip()
    if tail:
        parts.append(tail)
    return parts or [block]


def _extract_leading_title(block: str) -> tuple[str | None, str]:
    cleaned = _clean_paragraph(block)
    if not cleaned:
        return None, ""

    match = _LEADING_DATE_TITLE_RE.match(cleaned)
    if match:
        return _clean_heading(match.group("title")), match.group("rest").strip()

    if _looks_like_standalone_heading(cleaned):
        return _clean_heading(cleaned), ""

    return None, cleaned


def _looks_like_standalone_heading(block: str) -> bool:
    text = _clean_paragraph(block)
    if not (2 <= len(text) <= 140):
        return False
    if "\n" in block.strip():
        return False
    if (
        _CONTENTS_HEADING_RE.match(text)
        or _DECLARED_HEADING_RE.match(text)
        or _FRONT_BACK_HEADING_RE.match(text)
    ):
        return True
    if _LEADING_DATE_TITLE_RE.match(text):
        return True
    if re.match(r"^(?:[IVXLCDM]+|\d+)\.?$", text, re.IGNORECASE):
        return True
    if len(text) > 90 and text.endswith((".", "!", "?")):
        return False

    letters = re.findall(rf"[{_LETTER_RE_CHARS}]", text)
    uppercase = re.findall(rf"[{_UPPER_RE_CHARS}]", text)
    return bool(letters and len(uppercase) / len(letters) > 0.72)


def _split_chronicle_markers(paragraph: str) -> list[str]:
    if len(paragraph) < 900:
        return [paragraph]
    parts = [p.strip() for p in _CHRONICLE_MARKER_RE.split(paragraph) if p and p.strip()]
    return parts or [paragraph]


def _split_long_paragraph(paragraph: str, target: int = 950) -> list[str]:
    text = _clean_paragraph(paragraph)
    if len(text) <= 1200:
        return [text] if text else []

    sentences = [s.strip() for s in _SENTENCE_BOUNDARY_RE.split(text) if s.strip()]
    if len(sentences) <= 1:
        return _hard_wrap_paragraph(text, target)

    grouped: list[str] = []
    current: list[str] = []
    current_len = 0
    for sentence in sentences:
        sentence_len = len(sentence)
        if current and current_len + sentence_len > target and current_len >= 450:
            grouped.append(" ".join(current).strip())
            current = [sentence]
            current_len = sentence_len
        else:
            current.append(sentence)
            current_len += sentence_len + 1

    if current:
        tail = " ".join(current).strip()
        if grouped and len(tail) < 220:
            grouped[-1] = f"{grouped[-1]} {tail}"
        else:
            grouped.append(tail)
    return grouped


def _hard_wrap_paragraph(text: str, target: int) -> list[str]:
    parts: list[str] = []
    remaining = text
    while len(remaining) > target:
        split_at = remaining.rfind(" ", 0, target)
        if split_at < 450:
            split_at = target
        parts.append(remaining[:split_at].strip())
        remaining = remaining[split_at:].strip()
    if remaining:
        parts.append(remaining)
    return parts


def _body_paragraphs(block: str) -> list[str]:
    paragraphs: list[str] = []
    for candidate in _split_chronicle_markers(_clean_paragraph(block)):
        paragraphs.extend(_split_long_paragraph(candidate))
    return [p for p in paragraphs if p]


def _build_output_sections(text: str, document_title: str | None = None) -> list[_OutputSection]:
    plan = EditorialAssemblyAgent().assemble(text, document_title)
    return [
        _OutputSection(title=section.title, paragraphs=section.paragraphs)
        for section in plan.sections
    ] or [_OutputSection(title=_clean_heading(document_title or "Libro traducido"), paragraphs=[])]


def _chunk_untitled_sections(paragraphs: list[str], chunk_size: int = 55) -> list[_OutputSection]:
    chunks: list[_OutputSection] = []
    for idx in range(0, len(paragraphs), chunk_size):
        chunks.append(
            _OutputSection(
                title=f"Parte {len(chunks) + 1}",
                paragraphs=paragraphs[idx:idx + chunk_size],
            )
        )
    return chunks


def _paragraphs(text: str) -> list[str]:
    paras = [p.strip() for p in re.split(r"\n{2,}", text or "") if p.strip()]
    if not paras:
        paras = [line.strip() for line in (text or "").splitlines() if line.strip()]
    expanded: list[str] = []
    for para in paras:
        expanded.extend(_body_paragraphs(para))
    return expanded or [""]


def _write_docx(text: str, destination: Path) -> None:
    from docx import Document

    doc = Document()
    for para in _paragraphs(text):
        doc.add_paragraph(para)
    doc.save(str(destination))


def _write_pdf(text: str, destination: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import inch
    from xml.sax.saxutils import escape

    doc = SimpleDocTemplate(
        str(destination),
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )
    styles = getSampleStyleSheet()
    normal = styles["BodyText"]
    story = []
    for para in _paragraphs(text):
        story.append(Paragraph(escape(para), normal))
        story.append(Spacer(1, 0.12 * inch))
    doc.build(story)


def _write_epub(
    text: str,
    destination: Path,
    title: str | None = None,
    structure_hints: dict | None = None,
) -> None:
    book_title = _clean_heading(title or destination.stem or "Libro traducido")
    text, epub_tables = _extract_epub_markdown_tables(text)
    sections = _apply_epub_structure_hints(
        _build_output_sections(text, book_title),
        structure_hints,
    )
    sections = _maybe_sort_historical_anthology_sections(sections, book_title)
    book_id = f"urn:uuid:{uuid.uuid4()}"
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    chapter_files = [f"OEBPS/chap-{idx:03d}.xhtml" for idx in range(1, len(sections) + 1)]

    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destination, "w") as zf:
        mimetype_info = zipfile.ZipInfo("mimetype")
        mimetype_info.compress_type = zipfile.ZIP_STORED
        zf.writestr(mimetype_info, "application/epub+zip")

        zf.writestr(
            "META-INF/container.xml",
            _epub_container_xml(),
            compress_type=zipfile.ZIP_DEFLATED,
        )
        zf.writestr(
            "OEBPS/styles.css",
            _epub_styles_css(),
            compress_type=zipfile.ZIP_DEFLATED,
        )
        zf.writestr(
            "OEBPS/nav.xhtml",
            _epub_nav_xhtml(book_title, sections, chapter_files),
            compress_type=zipfile.ZIP_DEFLATED,
        )
        zf.writestr(
            "OEBPS/content.opf",
            _epub_package_opf(book_title, book_id, modified, chapter_files),
            compress_type=zipfile.ZIP_DEFLATED,
        )

        for idx, (section, filename) in enumerate(zip(sections, chapter_files), start=1):
            zf.writestr(
                filename,
                _epub_chapter_xhtml(section, idx, epub_tables),
                compress_type=zipfile.ZIP_DEFLATED,
            )


def _epub_container_xml() -> str:
    return """<?xml version="1.0" encoding="utf-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>
"""


def _apply_epub_structure_hints(
    sections: list[_OutputSection],
    structure_hints: dict | None,
) -> list[_OutputSection]:
    """Use optional LLM/rule hints to filter false heading candidates."""
    if not sections or not isinstance(structure_hints, dict):
        return sections

    raw_keep = structure_hints.get("keep_indices")
    if not isinstance(raw_keep, list):
        return sections

    keep_indices: set[int] = set()
    for value in raw_keep:
        try:
            keep_indices.add(int(value))
        except (TypeError, ValueError):
            continue

    if not keep_indices:
        return sections
    keep_indices.add(1)

    filtered: list[_OutputSection] = []
    for index, section in enumerate(sections, start=1):
        if index in keep_indices or not filtered:
            filtered.append(section)
            continue

        # Preserve every word: rejected headings become ordinary body text in
        # the previous section instead of disappearing from the export.
        filtered[-1].paragraphs.extend([section.title, *section.paragraphs])

    return [section for section in filtered if section.title or section.paragraphs]


def _maybe_sort_historical_anthology_sections(
    sections: list[_OutputSection],
    book_title: str,
) -> list[_OutputSection]:
    """Repair OCR EPUB order for dated historical anthologies only.

    Some OCR-built anthologies have spine chunks out of order while each entry
    title carries a date. Sorting ordinary fiction by dates would be dangerous,
    so this activates only when the filename/title strongly signals a history
    anthology and there are many dated sections.
    """
    if len(sections) < 40:
        return sections

    title_signal = _normalize_match_text(book_title)
    if not any(token in title_signal for token in ("history", "historia", "historico", "historica")):
        return sections

    intro: list[tuple[int, _OutputSection]] = []
    dated: list[tuple[int, int, _OutputSection]] = []
    undated: list[tuple[int, _OutputSection]] = []
    for index, section in enumerate(sections):
        if _normalize_match_text(section.title) in {"introduccion", "introduction"}:
            intro.append((index, section))
            continue
        year = _section_title_year(section)
        if year is None:
            undated.append((index, section))
        else:
            dated.append((year, index, section))

    if len(dated) < 40 or len(dated) / max(1, len(sections)) < 0.65:
        return sections

    sorted_sections = [section for _index, section in intro]
    sorted_sections.extend(
        section for _year, _index, section in sorted(dated, key=lambda item: (item[0], item[1]))
    )
    sorted_sections.extend(section for _index, section in undated)
    return sorted_sections


def _normalize_match_text(value: str) -> str:
    import unicodedata

    normalized = unicodedata.normalize("NFKD", value or "")
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", ascii_text)).strip()


def _section_title_year(section: _OutputSection) -> int | None:
    title = section.title or ""
    lower = title.lower()

    bc = re.search(r"(\d{1,4})\s*(?:a\.?\s*c\.?|a\s+de\s+c)", lower)
    if bc:
        return -int(bc.group(1))

    split_century = re.search(r"c\.\s*(\d{1,2})\.\s*(\d{2})\b", lower)
    if split_century:
        return int(split_century.group(1) + split_century.group(2))

    circa_year = re.search(r"c\.\s*(\d{3,4})\b", lower)
    if circa_year:
        year = int(circa_year.group(1))
        if 100 <= year <= 199 and ("aztec" in lower or "azteca" in lower):
            return year * 10
        return year

    full_years = [
        int(year)
        for year in re.findall(r"(?<!\d)(1[0-9]{3}|20[0-9]{2})(?!\d)", lower)
    ]
    if full_years:
        return full_years[0]

    ancient_bc_keywords = ("atenas", "griegos", "sócrates", "socrates")
    if any(keyword in lower for keyword in ancient_bc_keywords):
        ancient_years = [
            int(year)
            for year in re.findall(r"(?<!\d)([3-5][0-9]{2})(?!\d)", lower)
        ]
        if ancient_years:
            return -ancient_years[0]

    if "great eastern" in lower and section.paragraphs:
        paragraph_year = re.search(r"\b(1[0-9]{3})\b", section.paragraphs[0])
        if paragraph_year:
            return int(paragraph_year.group(1))

    return None


def _epub_styles_css() -> str:
    return """
@page {
  margin: 8%;
}
html {
  color-scheme: light;
}
body {
  font-family: Georgia, "Times New Roman", serif;
  line-height: 1.55;
  margin: 0;
  padding: 0;
  hyphens: auto;
  widows: 2;
  orphans: 2;
}
section.chapter {
  break-before: page;
  page-break-before: always;
}
h1 {
  font-size: 1.45em;
  line-height: 1.25;
  margin: 0 0 1.2em;
  text-align: left;
  break-after: avoid;
  page-break-after: avoid;
}
p {
  margin: 0 0 0.9em;
  text-indent: 1.1em;
  widows: 2;
  orphans: 2;
}
p.first,
p.formula {
  text-indent: 0;
}
p.first {
  margin-top: 0.2em;
}
p.formula {
  font-family: "Courier New", monospace;
  white-space: pre-wrap;
}
.table-wrap {
  margin: 1em 0 1.2em;
  overflow-x: auto;
}
table {
  border-collapse: collapse;
  font-size: 0.9em;
  line-height: 1.35;
  min-width: 100%;
}
th,
td {
  border: 1px solid #777;
  padding: 0.35em 0.5em;
  vertical-align: top;
}
th {
  font-weight: bold;
  background: #f2f2f2;
}
nav ol {
  list-style-type: none;
  padding-left: 0;
}
nav li {
  margin: 0.35em 0;
}
""".strip()


def _epub_nav_xhtml(
    book_title: str,
    sections: list[_OutputSection],
    chapter_files: list[str],
) -> str:
    items = []
    for section, filename in zip(sections, chapter_files):
        href = posixpath.basename(filename)
        items.append(
            f'<li><a href="{html.escape(href, quote=True)}">'
            f"{html.escape(section.title)}</a></li>"
        )
    return f"""<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="es" xml:lang="es">
  <head>
    <title>{html.escape(book_title)}</title>
    <link rel="stylesheet" type="text/css" href="styles.css"/>
  </head>
  <body>
    <nav epub:type="toc" id="toc">
      <h1>{html.escape(book_title)}</h1>
      <ol>
        {"".join(items)}
      </ol>
    </nav>
  </body>
</html>
"""


def _epub_package_opf(
    book_title: str,
    book_id: str,
    modified: str,
    chapter_files: list[str],
) -> str:
    manifest_items = [
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
        '<item id="css" href="styles.css" media-type="text/css"/>',
    ]
    spine_items = []
    for idx, filename in enumerate(chapter_files, start=1):
        href = posixpath.basename(filename)
        manifest_items.append(
            f'<item id="chap{idx}" href="{html.escape(href, quote=True)}" '
            'media-type="application/xhtml+xml"/>'
        )
        spine_items.append(f'<itemref idref="chap{idx}"/>')

    return f"""<?xml version="1.0" encoding="utf-8"?>
<package version="3.0" unique-identifier="bookid" xmlns="http://www.idpf.org/2007/opf" xmlns:dc="http://purl.org/dc/elements/1.1/">
  <metadata>
    <dc:identifier id="bookid">{html.escape(book_id)}</dc:identifier>
    <dc:title>{html.escape(book_title)}</dc:title>
    <dc:language>es</dc:language>
    <meta property="dcterms:modified">{html.escape(modified)}</meta>
  </metadata>
  <manifest>
    {"".join(manifest_items)}
  </manifest>
  <spine>
    {"".join(spine_items)}
  </spine>
</package>
"""


def _extract_epub_markdown_tables(text: str) -> tuple[str, dict[str, str]]:
    """Replace Markdown table blocks with stable placeholders for assembly."""
    table_map: dict[str, str] = {}
    blocks = re.split(r"(\n{2,})", text or "")
    rendered_blocks: list[str] = []
    table_index = 0

    def next_placeholder(lines: list[str]) -> str | None:
        nonlocal table_index
        table_html = _markdown_table_lines_to_xhtml(lines)
        if not table_html:
            return None
        placeholder = f"[[EPUB_TABLE_BLOCK_{table_index:04d}]]"
        table_map[placeholder] = table_html
        table_index += 1
        return placeholder

    for block in blocks:
        if not block or re.fullmatch(r"\n{2,}", block):
            rendered_blocks.append(block)
            continue
        lines = block.splitlines()
        output_lines: list[str] = []
        table_lines: list[str] = []
        for line in lines:
            stripped = line.strip()
            if stripped and _is_markdown_table_row(stripped):
                table_lines.append(stripped)
                continue
            if table_lines:
                placeholder = next_placeholder(table_lines)
                output_lines.extend([placeholder] if placeholder else table_lines)
                table_lines = []
            cleaned_line = _strip_malformed_markdown_table_artifact(line)
            if cleaned_line.strip():
                output_lines.append(cleaned_line)
        if table_lines:
            placeholder = next_placeholder(table_lines)
            output_lines.extend([placeholder] if placeholder else table_lines)
        rendered_blocks.append("\n".join(output_lines))
    return "".join(rendered_blocks), table_map


def _epub_chapter_xhtml(
    section: _OutputSection,
    chapter_index: int,
    table_map: dict[str, str] | None = None,
) -> str:
    body_parts = [f"<h1>{html.escape(section.title)}</h1>"]
    table_map = table_map or {}
    idx = 0
    visible_index = 0
    while idx < len(section.paragraphs):
        placeholder_html = _paragraph_table_placeholder_html(section.paragraphs[idx], table_map)
        if placeholder_html:
            body_parts.append(placeholder_html)
            idx += 1
            visible_index += 1
            continue
        table_html, next_idx = _consume_epub_markdown_table(section.paragraphs, idx)
        if table_html:
            body_parts.append(table_html)
            idx = next_idx
            visible_index += 1
            continue
        paragraph = section.paragraphs[idx]
        css_class = _paragraph_css_class(paragraph, is_first=visible_index == 0)
        class_attr = f' class="{css_class}"' if css_class else ""
        body_parts.append(f"<p{class_attr}>{html.escape(paragraph)}</p>")
        idx += 1
        visible_index += 1
    return f"""<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="es" xml:lang="es">
  <head>
    <title>{html.escape(section.title)}</title>
    <link rel="stylesheet" type="text/css" href="styles.css"/>
  </head>
  <body id="chapter-{chapter_index:03d}">
    <section class="chapter" epub:type="chapter">
      {"".join(body_parts)}
    </section>
  </body>
</html>
"""


def _paragraph_table_placeholder_html(paragraph: str, table_map: dict[str, str]) -> str:
    stripped = (paragraph or "").strip()
    if not stripped or not table_map:
        return ""
    if stripped in table_map:
        return table_map[stripped]
    matches = list(re.finditer(r"\[\[EPUB_TABLE_BLOCK_\d{4}\]\]", stripped))
    if not matches:
        return ""
    parts: list[str] = []
    cursor = 0
    for match in matches:
        before = stripped[cursor:match.start()].strip()
        if before:
            parts.append(f"<p>{html.escape(before)}</p>")
        parts.append(table_map.get(match.group(0), html.escape(match.group(0))))
        cursor = match.end()
    after = stripped[cursor:].strip()
    if after:
        parts.append(f"<p>{html.escape(after)}</p>")
    return "".join(parts)


def _consume_epub_markdown_table(paragraphs: list[str], start: int) -> tuple[str, int]:
    """Render a Markdown table block as native EPUB XHTML table markup."""
    if start >= len(paragraphs):
        return "", start

    first = paragraphs[start] or ""
    first_lines = [line.strip() for line in first.splitlines() if line.strip()]
    if first_lines and all(_is_markdown_table_row(line) for line in first_lines):
        html_table = _markdown_table_lines_to_xhtml(first_lines)
        if html_table:
            return html_table, start + 1

    if not _is_markdown_table_row(first.strip()):
        return "", start

    lines: list[str] = []
    idx = start
    while idx < len(paragraphs):
        raw = paragraphs[idx] or ""
        split_lines = [line.strip() for line in raw.splitlines() if line.strip()]
        if not split_lines or not all(_is_markdown_table_row(line) for line in split_lines):
            break
        lines.extend(split_lines)
        idx += 1

    html_table = _markdown_table_lines_to_xhtml(lines)
    if not html_table:
        return "", start
    return html_table, idx


def _is_markdown_table_row(line: str) -> bool:
    stripped = (line or "").strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def _strip_malformed_markdown_table_artifact(line: str) -> str:
    """Drop malformed Markdown table header fragments from prose output."""
    stripped = (line or "").strip()
    if not stripped or _is_markdown_table_row(stripped):
        return line
    if stripped.startswith("|") and _looks_like_table_fragment(stripped):
        return ""

    pipe_index = line.find("|")
    if pipe_index <= 0:
        return line
    tail = line[pipe_index:].strip()
    if tail.startswith("|") and _looks_like_table_fragment(tail):
        return line[:pipe_index].rstrip()
    return line


def _looks_like_table_fragment(value: str) -> bool:
    normalized = value.lower()
    if normalized.count("|") < 1:
        return False
    table_tokens = (
        "model",
        "modelo",
        "value",
        "valor",
        "bleu",
        "flop",
        "params",
        "ppl",
        "train",
        "entren",
        "en-de",
        "en-fr",
    )
    return any(token in normalized for token in table_tokens)


def _markdown_table_lines_to_xhtml(lines: list[str]) -> str:
    rows = [_split_markdown_table_row(line) for line in lines if _is_markdown_table_row(line)]
    rows = [row for row in rows if row]
    if len(rows) < 2:
        return ""

    has_header = len(rows) >= 2 and _is_markdown_separator_row(rows[1])
    if has_header:
        header = rows[0]
        body = rows[2:]
    else:
        header = []
        body = rows
    if not body and not header:
        return ""

    width = max(len(row) for row in ([header] if header else []) + body)

    def cells_html(row: list[str], tag: str) -> str:
        padded = row + [""] * (width - len(row))
        return "<tr>" + "".join(
            f"<{tag}>{html.escape(cell.strip())}</{tag}>"
            for cell in padded
        ) + "</tr>"

    parts = ['<div class="table-wrap"><table>']
    if header:
        parts.append("<thead>")
        parts.append(cells_html(header, "th"))
        parts.append("</thead>")
    if body:
        parts.append("<tbody>")
        for row in body:
            parts.append(cells_html(row, "td"))
        parts.append("</tbody>")
    parts.append("</table></div>")
    return "".join(parts)


def _split_markdown_table_row(line: str) -> list[str]:
    value = (line or "").strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]

    cells: list[str] = []
    current: list[str] = []
    escaped = False
    for char in value:
        if escaped:
            current.append(char)
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == "|":
            cells.append("".join(current).strip())
            current = []
            continue
        current.append(char)
    cells.append("".join(current).strip())
    return cells


def _is_markdown_separator_row(row: list[str]) -> bool:
    if not row:
        return False
    return all(re.fullmatch(r":?-{3,}:?", cell.strip() or "") for cell in row)


def _paragraph_css_class(paragraph: str, *, is_first: bool) -> str:
    classes: list[str] = []
    if is_first:
        classes.append("first")
    if _looks_like_formula(paragraph):
        classes.append("formula")
    return " ".join(classes)


def _looks_like_formula(paragraph: str) -> bool:
    if len(paragraph) > 320:
        return False
    return bool(
        re.search(r"(?:^|\s)[A-Za-z][A-Za-z0-9_]*\s*=", paragraph)
        or re.search(r"[=+\-*/^<>]\s*[A-Za-z0-9(]", paragraph)
        or re.search(r"\\(?:frac|sum|int|sqrt|alpha|beta|gamma)", paragraph)
        or re.search(r"[\u2211\u222B\u221A\u2264\u2265]", paragraph)
    )
