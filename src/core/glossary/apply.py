"""
Apply glossary corrections to already translated editable files.

The glossary controls future translation chunks. These helpers are for the
separate workflow where a user notices an already translated term and wants a
reviewed replacement propagated through an existing output document.
"""
from __future__ import annotations

import io
from pathlib import Path
import re
import zipfile
from typing import Dict, Iterable, List, Tuple


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".srt", ".html", ".htm", ".xhtml"}
EDITABLE_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".epub"}
UNSUPPORTED_EXTENSIONS = {".pdf"}


def _has_word_edge(text: str) -> bool:
    return bool(text) and (text[0].isalnum() or text[0] == "_" or text[-1].isalnum() or text[-1] == "_")


def _replacement_pattern(old: str, case_sensitive: bool) -> re.Pattern:
    escaped = re.escape(old)
    if _has_word_edge(old):
        escaped = rf"(?<!\w){escaped}(?!\w)"
    flags = 0 if case_sensitive else re.IGNORECASE
    return re.compile(escaped, flags)


def apply_term_correction_to_text(
    text: str,
    old_target: str,
    new_target: str,
    *,
    case_sensitive: bool = True,
) -> Tuple[str, int]:
    """Replace an old translated term with a corrected target.

    Latin-like terms use word boundaries, so replacing ``Adam`` will not alter
    ``Madam``. Terms that start/end with punctuation are treated literally.
    """
    old_target = old_target or ""
    new_target = new_target or ""
    if not text or not old_target or old_target == new_target:
        return text, 0
    pattern = _replacement_pattern(old_target, case_sensitive)
    return pattern.subn(new_target, text)


def _decode_bytes(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace")


def _unique_copy_path(path: Path) -> Path:
    base = path.with_name(f"{path.stem} (glossary corrected){path.suffix}")
    if not base.exists():
        return base
    for i in range(2, 1000):
        candidate = path.with_name(f"{path.stem} (glossary corrected {i}){path.suffix}")
        if not candidate.exists():
            return candidate
    return path.with_name(f"{path.stem} (glossary corrected {path.stat().st_mtime_ns}){path.suffix}")


def _write_text_result(
    path: Path,
    old_target: str,
    new_target: str,
    *,
    in_place: bool,
    case_sensitive: bool,
) -> Tuple[Path, int]:
    original = _decode_bytes(path.read_bytes())
    updated, count = apply_term_correction_to_text(
        original,
        old_target,
        new_target,
        case_sensitive=case_sensitive,
    )
    output_path = path if in_place else _unique_copy_path(path)
    if count > 0 or output_path != path:
        output_path.write_text(updated, encoding="utf-8")
    return output_path, count


def _iter_docx_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_in_docx_runs(paragraph, old_target: str, new_target: str, case_sensitive: bool) -> int:
    if not paragraph.runs:
        return 0
    original = "".join(run.text for run in paragraph.runs)
    updated, count = apply_term_correction_to_text(
        original,
        old_target,
        new_target,
        case_sensitive=case_sensitive,
    )
    if count <= 0:
        return 0
    paragraph.runs[0].text = updated
    for run in paragraph.runs[1:]:
        run.text = ""
    return count


def _write_docx_result(
    path: Path,
    old_target: str,
    new_target: str,
    *,
    in_place: bool,
    case_sensitive: bool,
) -> Tuple[Path, int]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to edit DOCX files") from exc

    doc = Document(str(path))
    total = 0
    for paragraph in _iter_docx_paragraphs(doc):
        total += _replace_in_docx_runs(paragraph, old_target, new_target, case_sensitive)

    output_path = path if in_place else _unique_copy_path(path)
    if total > 0 or output_path != path:
        doc.save(str(output_path))
    return output_path, total


def _write_epub_result(
    path: Path,
    old_target: str,
    new_target: str,
    *,
    in_place: bool,
    case_sensitive: bool,
) -> Tuple[Path, int]:
    output_path = path if in_place else _unique_copy_path(path)
    total = 0
    buffer = io.BytesIO()

    with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(buffer, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            lower = item.filename.lower()
            if lower.endswith((".xhtml", ".html", ".htm", ".opf", ".ncx")):
                text = _decode_bytes(data)
                updated, count = apply_term_correction_to_text(
                    text,
                    old_target,
                    new_target,
                    case_sensitive=case_sensitive,
                )
                if count:
                    data = updated.encode("utf-8")
                    total += count
            dst.writestr(item, data)

    if total > 0 or output_path != path:
        output_path.write_bytes(buffer.getvalue())
    return output_path, total


def apply_term_correction_to_file(
    path: Path,
    old_target: str,
    new_target: str,
    *,
    in_place: bool = False,
    case_sensitive: bool = True,
) -> Dict:
    """Apply a correction to one editable translated output file."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in UNSUPPORTED_EXTENSIONS:
        return {
            "filename": path.name,
            "supported": False,
            "error": "PDF files are not safely editable in place; regenerate from an editable output format.",
            "replacements": 0,
        }
    if ext not in EDITABLE_EXTENSIONS:
        return {
            "filename": path.name,
            "supported": False,
            "error": f"Unsupported file type: {ext or 'unknown'}",
            "replacements": 0,
        }

    if ext == ".docx":
        output_path, replacements = _write_docx_result(
            path,
            old_target,
            new_target,
            in_place=in_place,
            case_sensitive=case_sensitive,
        )
    elif ext == ".epub":
        output_path, replacements = _write_epub_result(
            path,
            old_target,
            new_target,
            in_place=in_place,
            case_sensitive=case_sensitive,
        )
    else:
        output_path, replacements = _write_text_result(
            path,
            old_target,
            new_target,
            in_place=in_place,
            case_sensitive=case_sensitive,
        )

    return {
        "filename": path.name,
        "output_filename": output_path.name,
        "supported": True,
        "in_place": output_path == path,
        "replacements": replacements,
    }


def apply_term_correction_to_files(
    paths: Iterable[Path],
    old_target: str,
    new_target: str,
    *,
    in_place: bool = False,
    case_sensitive: bool = True,
) -> List[Dict]:
    return [
        apply_term_correction_to_file(
            path,
            old_target,
            new_target,
            in_place=in_place,
            case_sensitive=case_sensitive,
        )
        for path in paths
    ]
